import time
from pathlib import Path

from app import practice
from app.config import MEDIA_DIR


def wait(client, job_id, timeout=60):
    t = time.time()
    while time.time() - t < timeout:
        j = client.get(f"/api/jobs/{job_id}").json()
        if j["status"] in ("done", "error"):
            return j
        time.sleep(0.2)
    raise AssertionError("job timeout")


def test_local_answer_check():
    q = {"choice_index": 2, "options": ["a", "b", "c"]}
    assert practice.local_check(q, "2") is True
    assert practice.local_check(q, 3) is False
    assert practice.local_check({"answer": "$1\\frac{5}{12}$"}, "17/12") is True
    assert practice.local_check({"answer": "$1\\frac{5}{12}$"}, "1又5/12") is True
    assert practice.local_check({"answer": "48 元"}, "48") is True
    assert practice.local_check({"answer": "48 元"}, "50") is False
    assert practice.local_check({"answer": "0.5"}, "1/2") is True
    assert practice.local_check({"answer": "25%"}, "0.25") is True
    assert practice.local_check({"qtype": "是非題", "answer": "○"}, "o") is True
    assert practice.local_check({"qtype": "是非題", "answer": "×"}, "○") is False
    assert practice.local_check({"answer": "光合作用"}, "光合作用") is True
    assert practice.local_check({"answer": "光合作用"}, "呼吸作用") is None  # needs AI


def test_status_and_me(client):
    s = client.get("/api/status").json()
    assert s["ai_ready"] and s["students"]
    me = client.get("/api/me").json()
    assert me["level"]["level"] == 1 and len(me["quests"]) == 3


def test_scan_exam_end_to_end(client, sample):
    with open(sample, "rb") as f:
        r = client.post("/api/exams", files=[("files", ("p1.jpg", f, "image/jpeg"))],
                        data={"title": "", "subject": "", "remove_handwriting": "true", "remove_color": "true"})
    assert r.status_code == 200, r.text
    j = wait(client, r.json()["job_id"])
    assert j["status"] == "done", j
    assert j["result"]["questions"] == 5
    e = client.get(f"/api/exams/{r.json()['exam_id']}").json()
    assert e["status"] == "ready" and e["subject"] == "數學"
    assert e["title"] == "六年級數學 第三單元 小考"
    qs = e["questions"]
    assert [q["number"] for q in qs] == ["1", "2", "3", "4", "5"]
    by = {q["number"]: q for q in qs}
    assert by["2"]["is_wrong"] == 1                       # teacher marked wrong
    assert by["4"]["is_marked"] == 1 and by["4"]["is_wrong"] == 1  # star + AI says wrong
    assert by["4"]["figure_path"] and (MEDIA_DIR / by["4"]["figure_path"]).exists()
    assert by["4"]["bar_model"]["enabled"]
    assert all(q["solved"] for q in qs) and by["1"]["choice_index"] == 2
    page = e["pages"][0]
    assert (MEDIA_DIR / page["scan_path"]).exists() and page["meta"]["warped"] and not page["clean_path"]  # no 還原卷
    # the synthetic photo has a yellow and a pink highlighter stripe
    assert len(page["meta"]["highlights"]) >= 2


def test_bank_filters_and_groups(client):
    wrong = client.get("/api/questions?flags=unmastered").json()
    assert wrong["total"] == 2
    marked = client.get("/api/questions?flags=marked").json()
    assert marked["total"] == 1
    any_mark = client.get("/api/questions?flags_any=highlighted,marked,wrong").json()
    assert any_mark["total"] >= 2
    groups = client.get("/api/questions/groups").json()
    assert {g["subject"] for g in groups} == {"數學"}


def test_answer_srs_and_rewards(client):
    q = client.get("/api/questions?flags=unmastered").json()["items"]
    q2 = next(x for x in q if x["number"] == "2")
    r = client.post(f"/api/questions/{q2['id']}/answer", json={"answer": "1", "mode": "review"}).json()
    assert r["correct"] is False and r["reward"]["xp"] == 3
    r = client.post(f"/api/questions/{q2['id']}/answer", json={"answer": "2", "mode": "review"}).json()
    assert r["correct"] is True and r["reward"]["xp"] == 15
    after = client.get(f"/api/questions/{q2['id']}").json()
    assert after["srs_box"] == 1 and after["next_review"] > time.strftime("%Y-%m-%d")
    q4 = next(x for x in q if x["number"] == "4")
    r = client.post(f"/api/questions/{q4['id']}/answer", json={"answer": "48", "mode": "practice"}).json()
    assert r["correct"] is True
    me = client.get("/api/me").json()
    assert me["xp"] > 0 and me["streak"] == 1 and me["active_today"]


def test_tutor_hint_chat(client):
    q = client.get("/api/questions?flags=marked").json()["items"][0]
    png = (Path(__file__).parent / "sample_exam.jpg").read_bytes()
    r = client.post(f"/api/questions/{q['id']}/tutor", files={"image": ("w.png", png, "image/png")})
    assert r.status_code == 200, r.text
    assert r.json()["result"] == "correct" and r.json()["reward"]["xp"] == 15
    h = client.post(f"/api/questions/{q['id']}/hint", json={"level": 1}).json()
    assert h["hint"] == "提示1" and h["cached"]           # stored hint ladder, no AI call
    h = client.post(f"/api/questions/{q['id']}/hint", json={"level": 2, "image": "data:image/png;base64,iVBORw0KGgo="}).json()
    assert h["hint"] == "AI 提示 2"                         # adaptive hint uses the drawing
    text = client.post(f"/api/questions/{q['id']}/chat", json={"message": "我不懂"}).text
    assert "一起想想" in text
    hist = client.get(f"/api/questions/{q['id']}/chat").json()
    assert [m["role"] for m in hist] == ["user", "assistant"]


def test_similar_generate_knowledge(client):
    q = client.get("/api/questions?flags=marked").json()["items"][0]
    j = wait(client, client.post(f"/api/questions/{q['id']}/similar", json={"n": 3}).json()["job_id"])
    assert len(j["result"]["question_ids"]) == 3
    detail = client.get(f"/api/questions/{q['id']}").json()
    assert len(detail["similar"]) == 3
    j = wait(client, client.post("/api/generate", json={"subject": "數學", "grade": 6, "unit": "分數的除法", "n": 5}).json()["job_id"])
    assert len(j["result"]["question_ids"]) == 5
    j = wait(client, client.post("/api/knowledge/generate", json={"subject": "數學", "grade": 6, "unit": "分數的除法"}).json()["job_id"])
    assert j["status"] == "done"
    k = client.get("/api/knowledge?subject=數學&unit=分數的除法&grade=6").json()
    assert k["content"]["formulas"][0]["name"] == "圓面積"
    units = client.get("/api/curriculum?subject=數學&grade=6").json()
    u = next(x for x in units if x["unit"] == "分數的除法")
    assert u["has_knowledge"] and u["questions"] == 5


def test_practice_modes(client):
    ch = client.get("/api/practice/challenge").json()
    assert 0 < len(ch["questions"]) <= 10
    r = client.post("/api/practice/challenge/finish", json={"correct": 8, "total": 10}).json()
    assert r["reward"]["xp"] == 30
    s = client.get("/api/units/practice?subject=數學&unit=分數的除法&grade=6").json()
    assert len(s) == 5
    r = client.post("/api/units/finish", json={"subject": "數學", "grade": 6, "unit": "分數的除法", "correct": 5, "total": 5}).json()
    assert r["stars"] == 3
    r = client.post("/api/practice/drill/finish", json={"kind": "frac_add", "correct": 21, "total": 25, "seconds": 60}).json()
    assert r["best"] == 21
    d = client.get("/api/practice/duel").json()
    assert d["bot"]["name"]
    r = client.post("/api/practice/duel/finish", json={"win": True, "score": 500, "bot_score": 300}).json()
    assert r["reward"]["xp"] == 25
    badges = {b["id"]: b for b in client.get("/api/me/badges").json()}
    assert badges["first_scan"]["earned_at"] and badges["unit_3star"]["earned_at"] and badges["drill_fast"]["earned_at"]


def test_quests_shop_league(client):
    quests = client.get("/api/me").json()["quests"]
    for q in quests:
        if q["done"]:
            r = client.post(f"/api/me/quests/{q['id']}/claim").json()
            assert r["coins"] == q["reward"]
    shop = client.get("/api/shop").json()
    cheap = min(shop["items"], key=lambda i: i["price"])
    r = client.post("/api/shop/buy", json={"item": cheap["id"]})
    assert r.status_code in (200, 400)
    assert client.post("/api/me/equip", json={"avatar": "🐲"}).status_code == 400  # not owned
    lg = client.get("/api/league").json()
    assert any(x["me"] for x in lg["board"]) and len(lg["board"]) == 10


def test_papers_print_word(client):
    ids = [q["id"] for q in client.get("/api/questions?flags=unmastered").json()["items"]]
    p = client.post("/api/papers", json={"title": "錯題卷", "question_ids": ids}).json()
    paper = client.get(f"/api/papers/{p['id']}").json()
    assert len(paper["questions"]) == len(ids)
    tok = client.post("/api/print/spec", json={"title": "錯題卷", "question_ids": ids,
                                               "config": {"orientation": "landscape", "columns": 2, "answers": "end", "explanations": True}}).json()["token"]
    spec = client.get(f"/api/print/spec/{tok}").json()
    assert spec["config"]["orientation"] == "landscape" and len(spec["questions"]) == len(ids)
    r = client.get(f"/api/print/docx/{tok}")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    # homework needs the parent PIN
    assert client.post("/api/papers", json={"title": "作業", "kind": "homework", "question_ids": ids}).status_code == 403


def test_region_select(client):
    exams = client.get("/api/exams").json()
    e = client.get(f"/api/exams/{exams[-1]['id']}").json()
    page = e["pages"][0]
    j = wait(client, client.post(f"/api/pages/{page['id']}/region", json={"box": [20, 600, 980, 700]}).json()["job_id"])
    assert j["status"] == "done", j
    assert client.post(f"/api/pages/{page['id']}/edits", json={"edits": []}).status_code in (404, 405)  # 手動修圖 is gone


def test_parent_settings_and_report(client):
    assert client.get("/api/settings").status_code == 403
    s = client.get("/api/settings", headers={"X-Parent-Pin": "0000"}).json()
    assert s["api_key_set"] and "…" in s["api_key"]
    j = wait(client, client.post("/api/report").json()["job_id"])
    assert j["result"]["summary"] == "整體不錯"
    st = client.get("/api/stats").json()
    assert st["report"]["report"]["plan"]
    # site password locks the API but keeps status reachable
    client.put("/api/settings", headers={"X-Parent-Pin": "0000"}, json={"site_password": "pw"})
    client.cookies.clear()
    assert client.get("/api/me").status_code == 401
    assert client.get("/api/status").json()["needs_login"]
    assert client.post("/api/login", json={"password": "wrong"}).status_code == 403
    assert client.post("/api/login", json={"password": "pw"}).status_code == 200
    assert client.get("/api/me").status_code == 401          # still needs a profile login
    child = next(p for p in client.get("/api/status").json()["profiles"] if p["role"] == "student")
    assert client.post("/api/auth/login", json={"id": child["id"], "pin": "1234"}).status_code == 200
    assert client.get("/api/me").status_code == 200
    client.put("/api/settings", headers={"X-Parent-Pin": "0000"}, json={"site_password": ""})


def test_visualize_and_wrong_paper_lesson(client):
    wrong = client.get("/api/questions?flags=unmastered").json()["items"]
    q = wrong[0]
    j = wait(client, client.post(f"/api/questions/{q['id']}/visualize").json()["job_id"])
    assert j["status"] == "done", j
    q2 = client.get(f"/api/questions/{q['id']}").json()
    assert q2["visuals"][0]["type"] == "fraction_bar" and q2["concept"]
    ids = [x["id"] for x in wrong]
    p = client.post("/api/papers", json={"title": "錯題卷", "question_ids": ids,
                                         "config": {"group_by_type": True, "lesson": True}}).json()
    j = wait(client, client.post(f"/api/papers/{p['id']}/lesson").json()["job_id"])
    assert j["status"] == "done", j
    paper = client.get(f"/api/papers/{p['id']}").json()
    assert paper["lesson"]["concepts"][0]["visual"]["type"] == "fraction_bar"
    tok = client.post("/api/print/spec", json={"title": "錯題卷", "question_ids": ids, "paper_id": p["id"],
                                               "config": {"group_by_type": True, "lesson": True}}).json()["token"]
    spec = client.get(f"/api/print/spec/{tok}").json()
    assert spec["lesson"]["concepts"]
    secs = [x["section"] for x in spec["questions"]]
    assert secs[0].startswith("一、") and all(x["number"] for x in spec["questions"])
    # each section restarts at 1
    for sec in set(secs):
        nums = [int(x["number"]) for x in spec["questions"] if x["section"] == sec]
        assert nums == list(range(1, len(nums) + 1))
    assert client.get(f"/api/print/docx/{tok}").status_code == 200


def test_shared_bank_personal_wrong(client):
    st = client.get("/api/status").json()
    me = st["user"]
    assert me["role"] == "student" and me["name"] == "姊姊"
    my_wrong = client.get("/api/questions?flags=unmastered").json()["total"]
    assert my_wrong > 0
    total = client.get("/api/questions").json()["total"]
    q = client.get("/api/questions?flags=unmastered").json()["items"][0]
    # wrong counts accumulate per child and keep the wrong answers
    before = client.get(f"/api/questions/{q['id']}").json()["wrong_count"]
    r = client.post(f"/api/questions/{q['id']}/answer", json={"answer": "999", "mode": "review"}).json()
    assert r["correct"] is False and r["wrong_count"] == before + 1
    after = client.get(f"/api/questions/{q['id']}").json()
    assert after["wrong_count"] == before + 1 and "999" in after["wrong_answers"]
    # parent creates a 3rd-grade sibling
    sis = client.post("/api/students", headers={"X-Parent-Pin": "0000"},
                      json={"name": "妹妹", "grade": 3, "pin": "5678"}).json()["id"]
    assert client.post("/api/auth/login", json={"id": sis, "pin": "0000"}).status_code == 403
    assert client.post("/api/auth/login", json={"id": sis, "pin": "5678"}).status_code == 200
    assert client.get("/api/me").json()["name"] == "妹妹"
    assert client.get("/api/exams").json() == []                          # exams stay personal
    assert client.get("/api/questions").json()["total"] == total          # bank is shared
    assert client.get("/api/questions?flags=unmastered").json()["total"] == 0   # wrong list is personal
    sq = client.get(f"/api/questions/{q['id']}").json()
    assert sq["wrong_count"] == 0 and sq["student_answer"] == "" and not sq["can_edit"]
    assert client.put(f"/api/questions/{q['id']}", json={"stem": "改掉"}).status_code == 403
    assert client.put(f"/api/questions/{q['id']}", json={"is_favorite": 1}).status_code == 200
    assert client.get(f"/api/questions/{q['id']}").json()["is_favorite"] == 1
    units = client.get("/api/curriculum?subject=數學&grade=3").json()
    assert any(u["unit"] == "同分母分數的加減" for u in units)
    # the older child's state is untouched
    parent = next(p for p in client.get("/api/status").json()["profiles"] if p["role"] == "parent")
    assert client.post("/api/auth/login", json={"id": parent["id"], "pin": "0000"}).status_code == 200
    seen = client.get(f"/api/questions/{q['id']}", headers={"X-Student-Id": str(me["id"])}).json()
    assert seen["wrong_count"] == before + 1 and seen["is_favorite"] == 0
    assert any(f["name"] == "姊姊" for f in client.get(f"/api/questions/{q['id']}").json()["family"])
    r = client.post("/api/questions/bulk", json={"ids": [q["id"]], "action": "add_to_wrong", "student_id": sis}).json()
    assert r["added"] == 1
    assert client.get("/api/questions?flags=unmastered", headers={"X-Student-Id": str(sis)}).json()["total"] == 1
    league = client.get("/api/league").json()["board"]
    assert not any(r["name"] == "爸爸" for r in league if not r["me"])
    assert client.post("/api/auth/login", json={"id": me["id"], "pin": "1234"}).status_code == 200


def test_reteach_variants(client):
    q = client.get("/api/questions?flags=unmastered&sort=wrong").json()["items"][0]
    j = wait(client, client.post(f"/api/questions/{q['id']}/reteach").json()["job_id"])
    assert j["status"] == "done", j
    d = client.get(f"/api/questions/{q['id']}").json()
    assert d["reteach"] and len(d["reteach"][0]["approaches"]) >= 2 and d["reteach"][0]["diagnosis"]


def test_nas_settings_parent_only(client):
    assert client.get("/api/nas").status_code == 403
    h = {"X-Parent-Pin": "0000"}
    r = client.put("/api/nas", headers=h, json={"path": r"\\192.0.2.1\Public\AI", "user": "u", "password": "p", "enabled": False})
    assert r.status_code == 200 and r.json()["password_set"] and r.json()["user"] == "u"
    assert "password" not in r.json()
    r = client.post("/api/nas/test", headers=h)   # TEST-NET address: unreachable → friendly error
    assert r.status_code == 400 and "NAS" in r.json()["detail"]


def test_parent_kid_features(client):
    h = {"X-Parent-Pin": "0000"}
    a = client.get("/api/analysis").json()
    assert a["units"] and "advice" in a["units"][0] and a["time"] and "week" in a
    m = client.get("/api/bank/matrix?subject=數學").json()
    assert m and set(m[0]["levels"]) <= {"basic", "advanced", "challenge"}
    q = client.get("/api/questions?flags=unmastered").json()["items"][0]
    assert q["level"] in ("basic", "advanced", "challenge")
    # levels 1/3/5 come from the stored 3-step ladder, 2/4 from AI, plus a worked example
    assert client.post(f"/api/questions/{q['id']}/hint", json={"level": 3}).json()["cached"]
    assert client.post(f"/api/questions/{q['id']}/hint", json={"level": 4}).json()["hint"] == "AI 提示 4"
    assert client.post(f"/api/questions/{q['id']}/hint", json={"example": True}).json()["worked_steps"]
    r = client.post(f"/api/questions/{q['id']}/steps", json={"text": "1-3/5=2/5\n120×2/5=40"}).json()
    assert r["result"] == "wrong" and r["attempt_id"]
    assert client.post(f"/api/attempts/{r['attempt_id']}/reason", json={"reason": "calc"}).status_code == 200
    client.post(f"/api/questions/{q['id']}/chat", json={"message": "為什麼要減？"})
    note = client.post(f"/api/questions/{q['id']}/chat/note").json()
    assert note["title"] and any(n["id"] == note["id"] for n in client.get("/api/notes").json())
    assert client.post(f"/api/questions/{q['id']}/help", json={"message": "看不懂"}).status_code == 200
    j = wait(client, client.post("/api/plan", headers=h, json={"n": 6}).json()["job_id"])
    assert j["status"] == "done", j
    paper = client.get(f"/api/papers/{j['result']['paper_id']}").json()
    assert len(paper["questions"]) == 6 and paper["config"]["lesson"]
    assert client.post(f"/api/questions/{q['id']}/coach").status_code == 403          # parents only
    assert client.post("/api/goals", json={"title": "段考", "exam_date": "2099-01-01"}).status_code == 403
    gid = client.post("/api/goals", headers=h, json={"title": "段考", "exam_date": "2099-01-01",
                                                     "units": [[q["subject"], q["unit"] or "分數的乘法"]]}).json()["id"]
    me = client.get("/api/me").json()
    assert me["goals"][0]["days_left"] > 0 and "minutes_today" in me
    rid = client.post("/api/rewards", headers=h, json={"title": "看卡通 30 分鐘", "cost": 1}).json()["id"]
    coins = me["coins"]
    x = client.post(f"/api/rewards/{rid}/redeem").json()["id"]
    assert client.get("/api/me").json()["coins"] == coins - 1
    assert client.post(f"/api/redemptions/{x}", headers=h, json={"status": "rejected"}).status_code == 200
    assert client.get("/api/me").json()["coins"] == coins          # refunded
    lv = client.get(f"/api/units/levels?subject=數學&unit=分數的除法").json()
    assert [l["level"] for l in lv] == ["basic", "advanced", "challenge"]
    client.delete(f"/api/goals/{gid}", headers=h)


def test_less_typing(client, sample):
    """v3: AI fills in what the child or parent used to type."""
    h = {"X-Parent-Pin": "0000"}
    q = next(x for x in client.get("/api/questions?flags=").json()["items"] if not x["choice_index"] and x["answer"])
    r = client.post(f"/api/questions/{q['id']}/answer", json={"answer": "72"}).json()
    assert not r["correct"] and r["diagnose_job"]
    d = wait(client, r["diagnose_job"])
    assert d["status"] == "done" and d["result"]["likely_cause"].startswith("你可能") and d["result"]["reason"] == "misread"
    # same wrong answer again: served from the shared cache, no new AI call
    from tests import fake_ai
    n = fake_ai.calls.count("diagnose")
    r2 = client.post(f"/api/questions/{q['id']}/answer", json={"answer": " 72 "}).json()
    assert wait(client, r2["diagnose_job"])["result"]["likely_cause"] and fake_ai.calls.count("diagnose") == n
    # one-tap help: the message writes itself
    m = client.post(f"/api/questions/{q['id']}/help", json={}).json()["message"]
    assert "72" in m and "AI 判斷" in m
    # exam goal: range suggested from recent scans and open wrong questions
    sug = client.get("/api/goals/suggest", headers=h).json()
    assert sug["title"] in ("期中考", "期末考") and sug["exam_date"] and sug["units"]
    # 再拍一頁: append a page to an existing exam
    e = client.get("/api/exams").json()[0]
    with open(sample, "rb") as f:
        j = client.post(f"/api/exams/{e['id']}/pages", files={"files": ("p2.jpg", f, "image/jpeg")}).json()["job_id"]
    assert wait(client, j)["status"] == "done"
    assert len(client.get(f"/api/exams/{e['id']}").json()["pages"]) == e["n_pages"] + 1


def test_tablet_pad_handoff(client):
    """用 iPad 寫: the computer opens a session, a tablet with no login writes, the computer grades."""
    from fastapi.testclient import TestClient
    from app.main import app
    q = client.get("/api/questions?flags=").json()["items"][0]
    assert client.post("/api/pads", json={"question_id": q["id"], "origin": "javascript:alert(1)"}).status_code == 400
    r = client.post("/api/pads", json={"question_id": q["id"], "origin": "https://raspberrypi.example.ts.net"}).json()
    assert r["url"].startswith("https://raspberrypi.example.ts.net/pad.html?t=") and "<svg" in r["qr"]
    tablet = TestClient(app)                                   # no cookie at all
    assert tablet.get("/api/exams").status_code == 401         # the token opens only the pad
    st = tablet.get(f"/api/pad/{r['token']}?v=-1&side=tablet").json()
    assert st["question"]["stem"] == q["stem"] and st["v"] == 0
    stroke = [{"color": "#111", "size": 3, "eraser": False, "pts": [[10, 10, .5], [50, 40, .5]]}]
    v = tablet.put(f"/api/pad/{r['token']}", json={"strokes": stroke, "w": 800, "h": 900}).json()["v"]
    mirror = client.get(f"/api/pad/{r['token']}?v=0").json()
    assert mirror["v"] == v and mirror["strokes"] == stroke and mirror["tablet_online"]
    assert "strokes" not in client.get(f"/api/pad/{r['token']}?v={v}").json()   # nothing new, nothing resent
    tablet.post(f"/api/pad/{r['token']}/submit")
    assert client.get(f"/api/pad/{r['token']}?v={v}").json()["status"] == "submit"
    assert tablet.post(f"/api/pad/{r['token']}/result", json={"result": "correct"}).status_code == 403
    client.post(f"/api/pad/{r['token']}/result", json={"result": "wrong", "diagnosis": "第 2 行算錯"})
    back = tablet.get(f"/api/pad/{r['token']}?v={v}&side=tablet").json()
    assert back["status"] == "graded" and back["result"]["diagnosis"] == "第 2 行算錯"
    assert client.get("/api/pads/active").json()["token"] == r["token"]
    assert tablet.get("/api/pad/not-a-real-token").status_code == 404
    client.delete(f"/api/pads/{r['token']}")
    assert tablet.get(f"/api/pad/{r['token']}").status_code == 404


def test_unit_prepare_parallel(client):
    """闖關 questions are written in parallel batches, once, before the tap."""
    from tests import fake_ai
    body = {"subject": "數學", "grade": 6, "unit": "質數與質因數分解", "level": "basic"}
    n0 = fake_ai.calls.count("generate")
    a = client.post("/api/units/prepare", json=body).json()
    b = client.post("/api/units/prepare", json=body).json()        # second tap joins, no new job
    res = wait(client, a["job_id"])
    assert a["job_id"] and (b["job_id"] == a["job_id"] or (b["job_id"] is None and b["have"] >= 5))
    assert res["status"] == "done" and len(res["result"]["question_ids"]) == 6
    assert fake_ai.calls.count("generate") - n0 == 3                # 3 batches of 2
    assert client.post("/api/units/prepare", json=body).json()["job_id"] is None   # enough now
    s = client.get("/api/units/practice?subject=數學&unit=質數與質因數分解&grade=6&n=5&level=basic").json()
    assert len(s) == 5


def test_phone_photos_to_desktop(client, sample):
    """用手機拍: a phone with no login sends photos; the computer gets one exam."""
    from fastapi.testclient import TestClient
    from app.main import app
    r = client.post("/api/photo-sessions", json={"origin": "http://192.168.1.20:8000"}).json()
    assert r["url"].endswith(f"/snap.html?t={r['token']}") and "<svg" in r["qr"]
    phone = TestClient(app)
    assert phone.get("/api/exams").status_code == 401
    assert phone.get(f"/api/pad/{r['token']}").status_code == 404          # a photo token is not a pen token
    for _ in range(3):
        with open(sample, "rb") as f:
            n = phone.post(f"/api/snap/{r['token']}/photo", files={"file": ("p.jpg", f, "image/jpeg")}).json()["count"]
    assert n == 3
    # the camera sends each photo at once and fixes it up after its check (by id: positions move)
    with open(sample, "rb") as f:
        j = phone.post(f"/api/snap/{r['token']}/photo", files={"file": ("p.jpg", f, "image/jpeg")}).json()
    assert j["count"] == 4 and j["id"]
    assert phone.put(f"/api/snap/{r['token']}/p/{j['id']}", json={"corners": [[0.1, 0.1], [0.9, 0.1], [0.9, 0.9], [0.1, 0.9]], "spread": True}).json()["ok"]
    assert phone.get(f"/api/snap/{r['token']}/corners/3").json()["set"] == [[0.1, 0.1], [0.9, 0.1], [0.9, 0.9], [0.1, 0.9]]
    assert phone.put(f"/api/snap/{r['token']}/p/{j['id']}", json={"corners": "x"}).status_code == 400
    assert phone.delete(f"/api/snap/{r['token']}/p/{j['id']}").json()["count"] == 3
    assert phone.delete(f"/api/snap/{r['token']}/p/{j['id']}").status_code == 404
    assert phone.delete(f"/api/snap/{r['token']}/photo/1").json()["count"] == 2
    assert client.get(f"/api/snap/{r['token']}/thumb/0").status_code == 200
    st = client.get(f"/api/snap/{r['token']}").json()
    assert st["count"] == 2 and st["sent"] == [] and st["owner"]
    # 完成: one exam goes to the AI in the background and the same connection is ready for the
    # next exam at once (the test's phone client has no long-lived event loop for the
    # background job, so the computer finishes here)
    d = client.post(f"/api/snap/{r['token']}/done").json()
    eid = d["exam_id"]
    assert d["n"] == 1
    assert phone.post(f"/api/snap/{r['token']}/done").status_code == 400          # nothing new to send
    with open(sample, "rb") as f:
        assert phone.post(f"/api/snap/{r['token']}/photo", files={"file": ("p.jpg", f, "image/jpeg")}).json()["count"] == 1
    st = client.get(f"/api/snap/{r['token']}").json()
    assert st["count"] == 1 and [x["exam_id"] for x in st["sent"]] == [eid]
    job = client.get(f"/api/exams/{eid}").json()["job"]
    assert wait(client, job["id"])["status"] == "done"
    assert len(client.get(f"/api/exams/{eid}").json()["pages"]) == 2


def test_split_and_merge(monkeypatch):
    """Long AI work is split into parallel requests and merged (real code, fake transport)."""
    import asyncio
    from app import ai, prompts
    from tests import fake_ai
    seen = []

    async def fake_call(task, content, schema, effort="high", max_tokens=0, system_extra="", fast=False):
        text = content[-1]["text"]
        seen.append((task, fast))
        if task == "reteach":
            style = next(s for s in prompts.RETEACH_STYLES if s in text)
            return {"diagnosis": "可能把用掉的當剩下", "approaches": [{"name": style[:6], "why": "", "explanation": "", "steps": [], "visual": {}}],
                    "check_question": {"question": "", "answer": ""}, "prerequisite": {"unit": "", "reason": ""}, "encouragement": "加油"}
        if task == "lesson":
            nums = [int(n) for n in __import__("re").findall(r"第 (\d+) 題", text)]
            return {"title": "觀念", "intro": "先看懂", "concepts": [{"title": f"觀念{nums[0]}", "question_numbers": nums}], "tips": ["看清楚單位"]}
        if task == "similar":
            return {"questions": [{"stem": text[-20:]}]}
        raise AssertionError(task)

    real = {n: getattr(fake_ai.ai, n) for n in ("reteach", "lesson", "similar")}
    import importlib
    fresh = importlib.reload(importlib.import_module("app.ai"))   # the real functions, not the fakes
    monkeypatch.setattr(fresh, "call_json", fake_call)
    q = {"stem": "120 元用了 3/5", "answer": "48", "wrong_count": 3, "unit": "分數"}
    r = asyncio.run(fresh.reteach(q, [], 3))
    assert len(r["approaches"]) == 3 and r["diagnosis"] and all(f for t, f in seen if t == "reteach")
    items = [{"n": i + 1, "stem": f"題{i}", "unit": u, "qtype": "應用題"} for i, u in enumerate(["分數", "分數", "圓", "比", "速率", "體積"])]
    les = asyncio.run(fresh.lesson(items))
    assert [c["question_numbers"][0] for c in les["concepts"]] == sorted(c["question_numbers"][0] for c in les["concepts"])
    assert sum(t == "lesson" for t, _ in seen) == 4 and les["tips"] == ["看清楚單位"]          # 5 units folded into 4 requests
    assert len(asyncio.run(fresh.similar(q, 3))) == 3
    fake_ai.install()                                              # put the fakes back for the other tests
    for n, f in real.items():
        assert getattr(fake_ai.ai, n)


def test_prepared_before_asking(client):
    """Wrong twice → diagram and re-teach start; a hint → the next level and the example get ready."""
    from tests import fake_ai
    q = next(x for x in client.get("/api/questions?flags=").json()["items"] if not x["choice_index"] and x["answer"])
    for _ in range(2):
        client.post(f"/api/questions/{q['id']}/answer", json={"answer": "999"})
    time.sleep(0.5)
    from app import db
    kinds = {r["kind"] for r in db.q("SELECT kind FROM jobs WHERE ref_id=?", (q["id"],))}
    assert {"reteach", "visualize"} <= kinds or db.scalar("SELECT 1 FROM explanations WHERE question_id=? AND kind='reteach'", (q["id"],))
    n_hint = fake_ai.calls.count("hint")
    client.post(f"/api/questions/{q['id']}/hint", json={"level": 1})
    time.sleep(0.5)
    assert fake_ai.calls.count("hint") == n_hint + 1 and "example_hint" in fake_ai.calls   # level 2 + example prepared
    r = client.post(f"/api/questions/{q['id']}/hint", json={"level": 2}).json()
    assert r.get("cached") and fake_ai.calls.count("hint") == n_hint + 1                     # served from the cache
    assert client.post(f"/api/questions/{q['id']}/hint", json={"example": True}).json().get("cached")


def test_figure_correction(client):
    """「AI 看錯圖了」: the note is saved, the question re-solved, stale help removed."""
    from app import db
    q = client.get("/api/questions?flags=").json()["items"][0]
    db.execute("UPDATE questions SET figure_unsure='11 量到哪裡？' WHERE id=?", (q["id"],))
    db.insert("explanations", {"question_id": q["id"], "student_id": 0, "wrong_count": 2, "content": {"hint": "舊的"},
                               "kind": "hint", "created_at": db.now()})
    j = client.post(f"/api/questions/{q['id']}/solve", json={"note": "11 是 A 到最底下虛線的整段"}).json()["job_id"]
    assert wait(client, j)["status"] == "done"
    q2 = client.get(f"/api/questions/{q['id']}").json()
    assert q2["fix_note"].startswith("11 是") and q2["figure_unsure"] == ""
    assert not db.scalar("SELECT 1 FROM explanations WHERE question_id=? AND kind='hint'", (q["id"],))
    from app import ai
    assert "家長的更正說明" in ai.question_text(q2) and "參考解法也可能看錯圖" in ai.reference({**q2, "answer": "80"})


def test_visual_check_redraws_or_drops(client):
    """每張圖解畫完都要給 AI 看過：不清楚就照意見重畫，還是不行就不給孩子看。"""
    import asyncio

    from app import pipeline, vischeck
    from tests import fake_ai
    good, v = asyncio.run(vischeck.refine("題目", dict(fake_ai.VIS)))
    assert good and v["ok"]
    fixed, v = asyncio.run(vischeck.refine("題目", {**fake_ai.VIS, "title": "亂畫的圖"}))
    assert fixed["title"] == "重畫的分數條" and v["ok"]
    gone, v = asyncio.run(vischeck.refine("題目", {**fake_ai.VIS, "title": "亂畫壞到底"}))
    assert gone is None and v["problems"]

    # existing diagrams: the recheck keeps good ones (marked checked) and removes hopeless ones
    q = client.get("/api/questions").json()["items"][0]
    from app import db
    db.update("questions", q["id"], {"visuals": [dict(fake_ai.VIS), {**fake_ai.VIS, "title": "亂畫壞到底"}]})
    n = asyncio.run(pipeline.check_question(q["id"]))
    vis = client.get(f"/api/questions/{q['id']}").json()["visuals"]
    assert n == 2 and len(vis) == 1 and vis[0]["checked"]
    assert asyncio.run(pipeline.check_question(q["id"])) == 0      # already checked: not looked at again


def test_lesson_diagrams_appear_after_check(client):
    p = client.post("/api/papers", json={"title": "圖解檢查卷", "question_ids": [q["id"] for q in client.get("/api/questions").json()["items"][:2]],
                                         "config": {"lesson": True}}).json()
    wait(client, client.post(f"/api/papers/{p['id']}/lesson").json()["job_id"])
    import time
    for _ in range(50):
        c = client.get(f"/api/papers/{p['id']}").json()["lesson"]["concepts"][0]["visual"]
        if c.get("checked"):
            break
        time.sleep(0.1)
    assert c.get("checked") and not c.get("pending")


def test_templates_preview_download_upload(client):
    """試卷範本: every template previews, downloads as Word, and an uploaded Word
    template is filled with the questions exactly where its {{…}} marks are."""
    import io

    from docx import Document
    lst = client.get("/api/templates").json()
    ids = [t["id"] for t in lst["templates"]]
    assert {"quiz", "midterm", "mock", "private", "wrong", "unit"} <= set(ids)
    exam = client.get("/api/exams").json()[0]
    for tid in ids:
        tok = client.post(f"/api/templates/{tid}/sample", json={"exam_id": exam["id"]}).json()["token"]
        spec = client.get(f"/api/print/spec/{tok}").json()
        assert spec["layout"]["template"] == tid and spec["questions"]
        assert client.get(f"/api/print/docx/{tok}").status_code == 200
    mid = client.get(f"/api/print/spec/{client.post('/api/templates/midterm/sample', json={}).json()['token']}").json()
    assert "學年度" in mid["title"] and mid["layout"]["score_table"] and "分）" in mid["questions"][0]["section"]
    # download → "edit in Word" → upload
    raw = client.get("/api/templates/midterm/docx").content
    doc = Document(io.BytesIO(raw))
    doc.paragraphs[1].insert_paragraph_before("台北市立快樂國小")
    buf = io.BytesIO(); doc.save(buf)
    P = {"x-parent-pin": "0000"}
    r = client.post("/api/templates", files={"file": ("我的段考.docx", buf.getvalue())}, data={"name": "我的段考"}, headers=P)
    assert r.status_code == 200, r.text
    meta = r.json()
    assert meta["base"] == "midterm"
    bad = client.post("/api/templates", files={"file": ("x.docx", b"not a docx")}, data={"name": "x"}, headers=P)
    assert bad.status_code == 400
    tok = client.post(f"/api/templates/{meta['id']}/sample", json={"exam_id": exam["id"]}).json()["token"]
    out = Document(io.BytesIO(client.get(f"/api/print/docx/{tok}").content))
    text = "\n".join(p.text for p in out.paragraphs)
    assert "台北市立快樂國小" in text and "{{" not in text and "【範本說明】" not in text
    qs = client.get(f"/api/print/spec/{tok}").json()["questions"]
    assert qs[0]["stem"].split("\n")[0][:6].replace("$", "") in text.replace("$", "") or len(text) > 50
    client.put("/api/templates/default", headers={"X-Parent-Pin": "0000"}, json={"id": meta["id"]})
    assert client.get("/api/templates").json()["default"] == meta["id"]
    client.delete(f"/api/templates/{meta['id']}", headers=P)
    assert meta["id"] not in [t["id"] for t in client.get("/api/templates").json()["templates"]]


def test_bank_overview_and_tasks(client):
    """錯題頁: concepts to work on (child) / weaknesses + tasks (parent)."""
    from app import db
    qs_ = client.get("/api/questions?flags=wrong&limit=50").json()["items"]
    assert qs_, "the sample exam has wrong questions"
    ov = client.get("/api/bank/overview").json()
    assert ov["stats"]["wrong"] >= 1 and "concepts" in ov and "picks" in ov
    assert "weaknesses" not in ov                                   # a child does not get the parent part
    c = ov["concepts"][0]
    assert c["ids"] and c["open"] == len(c["ids"]) and c["unit"]
    r = client.get("/api/questions", params={"ids": ",".join(map(str, c["ids"]))}).json()
    assert {q["id"] for q in r["items"]} == set(c["ids"])
    P = {"x-parent-pin": "0000"}
    assert client.post("/api/tasks", json={"units": [[c["subject"], c["unit"]]], "kind": "redo"}).status_code == 403
    t = client.post("/api/tasks", json={"units": [[c["subject"], c["unit"]]], "kind": "redo", "n": 5, "reward": 20,
                                        "message": "加油"}, headers=P)
    assert t.status_code == 200, t.text
    pid = t.json()["id"]
    me = client.get("/api/me").json()
    assert any(h["id"] == pid for h in me["homework"])                # shows on her 今天
    paper = client.get(f"/api/papers/{pid}").json()
    assert paper["config"]["task"]["kind"] == "redo" and paper["questions"]
    for q in paper["questions"]:
        client.post(f"/api/questions/{q['id']}/answer", json={"answer": q["answer"], "mode": "review"})
    tasks = [x for x in client.get("/api/bank/overview", headers=P).json()["tasks"] if x["id"] == pid]
    assert tasks and tasks[0]["done"] == len(paper["questions"])
    done = client.post(f"/api/papers/{pid}/complete").json()
    assert done["reward"]
    assert client.delete(f"/api/tasks/{pid}", headers=P).status_code == 200
    db.execute("DELETE FROM papers WHERE id=?", (pid,))


def test_exam_template_and_kind_mapping(client):
    """考卷頁: what the exam turns into; 這類考卷都用我的 Word 範本."""
    import io

    from docx import Document
    P = {"x-parent-pin": "0000"}
    exam = client.get("/api/exams").json()[0]
    t = client.get(f"/api/exams/{exam['id']}").json()["template"]
    assert t["id"] and t["kind"] and t["auto"]
    r = client.put(f"/api/exams/{exam['id']}/template", json={"id": "mock"}).json()
    assert r["id"] == "mock" and not r["auto"]
    # upload my Word for this exam's kind and map the kind to it
    client.put(f"/api/exams/{exam['id']}/template", json={"id": "auto"})
    kind = client.get(f"/api/exams/{exam['id']}").json()["template"]["kind"]
    doc = Document(io.BytesIO(client.get(f"/api/templates/{kind}/docx").content))
    buf = io.BytesIO(); doc.save(buf)
    mine = client.post("/api/templates", files={"file": ("mine.docx", buf.getvalue())}, data={"name": "我的"}, headers=P).json()
    assert client.put("/api/templates/for", json={"kind": kind, "id": mine["id"]}).status_code == 403
    m = client.put("/api/templates/for", json={"kind": kind, "id": mine["id"]}, headers=P).json()["mapping"]
    assert m[kind] == mine["id"]
    t = client.get(f"/api/exams/{exam['id']}").json()["template"]
    assert t["id"] == mine["id"] and t["custom"] and t["auto"]
    tok = client.post(f"/api/templates/{t['id']}/sample", json={"exam_id": exam["id"], "key_only": True, "answer_style": "steps"}).json()["token"]
    assert client.get(f"/api/print/spec/{tok}").json()["config"]["key_only"]
    assert client.get(f"/api/print/docx/{tok}").status_code == 200
    client.put("/api/templates/for", json={"kind": kind, "id": ""}, headers=P)
    client.delete(f"/api/templates/{mine['id']}", headers=P)
    assert client.get(f"/api/exams/{exam['id']}").json()["template"]["id"] == kind


def test_scan_owner_and_move(client, sample):
    """考卷要放進孩子的錯題本：家長帳號掃描時歸給孩子；放錯人可以改。"""
    from fastapi.testclient import TestClient

    from app import db, main
    kid = db.scalar("SELECT id FROM students WHERE role='student' ORDER BY id LIMIT 1")
    with TestClient(main.app) as pc:
        parent = next(p for p in pc.get("/api/status").json()["profiles"] if p["role"] == "parent")
        assert pc.post("/api/auth/login", json={"id": parent["id"], "pin": "0000"}).status_code == 200
        kids = db.q("SELECT id FROM students WHERE role='student'")
        with open(sample, "rb") as f:
            r = pc.post("/api/exams", files={"files": ("p.jpg", f, "image/jpeg")}, data={"title": "家長拍的"})
        if len(kids) > 1:   # several children: the parent must say whose it is
            assert r.status_code == 400 and "哪個孩子" in r.json()["detail"]
            with open(sample, "rb") as f:
                r = pc.post("/api/exams", files={"files": ("p.jpg", f, "image/jpeg")}, data={"title": "家長拍的", "student_id": kid})
        assert r.status_code == 200, r.text
        eid = r.json()["exam_id"]
        assert db.scalar("SELECT student_id FROM exams WHERE id=?", (eid,)) == kid     # the child's, not the parent's
        job = pc.get(f"/api/exams/{eid}", headers={"x-student-id": str(kid)}).json()["job"]
        wait(pc, job["id"])
        # an exam filed under the parent by mistake is moved to the child, marks and all
        old = eid
        qids = [x["id"] for x in db.q("SELECT id FROM questions WHERE exam_id=?", (old,))]
        db.execute("UPDATE exams SET student_id=? WHERE id=?", (parent["id"], old))
        db.execute(f"UPDATE student_q SET student_id=? WHERE student_id=? AND question_id IN ({','.join('?' * len(qids))})", [parent["id"], kid] + qids)
        r = pc.put(f"/api/exams/{old}/owner", json={"student_id": kid})
        assert r.status_code == 200, r.text
        assert db.scalar("SELECT student_id FROM exams WHERE id=?", (old,)) == kid
        assert db.scalar(f"SELECT COUNT(*) FROM student_q WHERE student_id=? AND question_id IN ({','.join('?' * len(qids))})", [kid] + qids) == len(qids)
        assert pc.put(f"/api/exams/{old}/owner", json={"student_id": parent["id"]}).status_code == 400   # never to a parent


def test_custom_exam_types(client):
    """整理清單: own types with keywords, hidden types, used by AI 自動判斷 and the exam page."""
    P = {"x-parent-pin": "0000"}
    assert client.post("/api/templates/types", json={"name": "週考"}).status_code == 403
    t = client.post("/api/templates/types", json={"name": "補習班週考", "base": "private", "keywords": ["週考"],
                                                  "config": {"answer_style": "brief", "columns": 1, "evil": 1}}, headers=P).json()
    lst = client.get("/api/templates").json()["templates"]
    mine = next(x for x in lst if x["id"] == t["id"])
    assert mine["kind"] == "type" and mine["config"]["answer_style"] == "brief" and mine["config"]["columns"] == 1 and "evil" not in mine["config"]
    exam = client.get("/api/exams").json()[0]
    old = exam["title"]
    client.put(f"/api/exams/{exam['id']}", json={"title": "第三次週考"})
    client.put(f"/api/exams/{exam['id']}/template", json={"id": "auto"})
    et = client.get(f"/api/exams/{exam['id']}").json()["template"]
    assert et["kind"] == t["id"] and et["kind_name"] == "補習班週考"
    assert client.put("/api/templates/unit/hidden", json={"hidden": True}, headers=P).status_code == 200
    assert next(x for x in client.get("/api/templates").json()["templates"] if x["id"] == "unit")["hidden"]
    client.put("/api/templates/unit/hidden", json={"hidden": False}, headers=P)
    client.delete(f"/api/templates/types/{t['id']}", headers=P)
    client.put(f"/api/exams/{exam['id']}", json={"title": old})
    assert t["id"] not in [x["id"] for x in client.get("/api/templates").json()["templates"]]


def test_regenerate_by_template(client):
    P = {"x-parent-pin": "0000"}
    exam = next(e for e in client.get("/api/exams").json() if e["status"] == "ready")
    client.put(f"/api/exams/{exam['id']}/template", json={"id": "mock"})
    dry = client.post("/api/exams/regenerate", json={"ids": [exam["id"]], "reset": True, "dry": True}, headers=P).json()["plan"]
    assert dry[0]["auto"] and client.get(f"/api/exams/{exam['id']}").json()["template"]["id"] == "mock"   # dry changes nothing
    r = client.post("/api/exams/regenerate", json={"ids": [exam["id"]], "reset": True}, headers=P).json()
    assert wait(client, r["job_id"])["status"] == "done"
    assert client.get(f"/api/exams/{exam['id']}").json()["template"]["auto"]


def test_job_starting_endpoints(client):
    """Endpoints that start background work must run on the event loop (a plain def failed)."""
    P = {"x-parent-pin": "0000"}
    ov = client.get("/api/bank/overview").json()
    c = ov["concepts"][0]
    r = client.post("/api/bank/similar", json={"ids": c["ids"][:2]})
    assert r.status_code == 200 and wait(client, r.json()["job_id"])["status"] == "done"
    for kind in ("lesson", "similar"):
        t = client.post("/api/tasks", json={"units": [[c["subject"], c["unit"]]], "kind": kind, "n": 3}, headers=P)
        assert t.status_code == 200, t.text
        client.delete(f"/api/tasks/{t.json()['id']}", headers=P)
    assert client.post("/api/visuals/recheck", json={}, headers=P).status_code == 200


def test_type_from_name_only(client):
    """只打名稱就能新增類型：排版自動選、關鍵字自動抓；AI 判斷的類型會被採用。"""
    from app import db, templates
    P = {"x-parent-pin": "0000"}
    t = client.post("/api/templates/types", json={"name": "學校段考"}, headers=P).json()
    assert t["base"] == "midterm" and t["keywords"] == ["段考"]
    w = client.post("/api/templates/types", json={"name": "誠明班"}, headers=P).json()
    assert w["base"] == "quiz" and w["keywords"] == ["誠明班"]
    exam = db.one("SELECT * FROM exams WHERE status='ready' ORDER BY id LIMIT 1")
    opts = exam["options"] if isinstance(exam["options"], dict) else {}
    db.update("exams", exam["id"], {"options": {**opts, "template": "auto", "kind_ai": w["id"]}})
    assert templates.exam_category(db.one("SELECT * FROM exams WHERE id=?", (exam["id"],))) == w["id"]
    db.update("exams", exam["id"], {"options": opts})
    for x in (t, w):
        client.delete(f"/api/templates/types/{x['id']}", headers=P)


def test_question_upgrade_and_layout(client):
    """題目升級 + 出題軟體排版: proofread, SVG / table figures, option notes; Word and print use them."""
    import io

    from docx import Document

    from app import db
    P = {"x-parent-pin": "0000"}
    exam = next(e for e in client.get("/api/exams").json() if e["status"] == "ready")
    r = client.post("/api/exams/upgrade", json={"exam_id": exam["id"], "force": True}, headers=P).json()
    assert wait(client, r["job_id"])["status"] == "done"
    qs = db.q("SELECT * FROM questions WHERE exam_id=?", (exam["id"],))
    assert all(q["upgraded_at"] for q in qs)
    q = qs[0]
    db.update("questions", q["id"], {"figure_table": {"header": ["", "甲", "乙"], "rows": [["流速", "24", "15"]]},
                                     "options": ["甲", "乙", "丙", "丁"], "qtype": "選擇題"})
    tok = client.post(f"/api/templates/private/sample", json={"exam_id": exam["id"]}).json()["token"]
    spec = client.get(f"/api/print/spec/{tok}").json()
    assert spec["config"]["columns"] == 1
    doc = Document(io.BytesIO(client.get(f"/api/print/docx/{tok}").content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "（　　）" in text and "①" in text and "\t" in text          # bracket + options side by side
    assert any("流速" in c.text for t in doc.tables for row in t.rows for c in row.cells)   # a real table
    key = client.post(f"/api/templates/private/sample", json={"exam_id": exam["id"], "key_only": True}).json()["token"]
    kdoc = Document(io.BytesIO(client.get(f"/api/print/docx/{key}").content))
    assert "① 對" in "\n".join(p.text for p in kdoc.paragraphs)


def test_upgrade_priority_order(client):
    """優先升級: the prioritized exam's questions are taken first by the parallel sessions."""
    from app import db, pipeline
    exams = [e["id"] for e in db.q("SELECT id FROM exams WHERE status='ready' ORDER BY id")]
    if len(exams) < 2:
        return
    old, new = exams[0], exams[-1]
    db.execute("UPDATE questions SET upgraded_at='' WHERE exam_id IN (?, ?)", (old, new))
    db.set_setting("upgrade_priority", [])
    first = db.scalar("SELECT exam_id FROM questions WHERE id=?", (pipeline._next_question(),))
    assert first == new                                    # newest first by default
    pipeline.prioritize(old)
    assert db.scalar("SELECT exam_id FROM questions WHERE id=?", (pipeline._next_question(),)) == old
    db.execute("UPDATE questions SET upgraded_at='test' WHERE exam_id IN (?, ?)", (old, new))
    db.set_setting("upgrade_priority", [])


def test_ai_lanes_user_first():
    """使用者操作優先: a person's AI call starts at once even when background work fills its lane."""
    import asyncio

    from app import ai, db
    db.set_setting("upgrade_parallel", 2)
    order = []

    async def call(name, bg, hold):
        async def run():
            if bg:
                ai.background.set(True)
            async with ai._lanes.slot():
                order.append(name)
                await asyncio.sleep(hold)
        return await asyncio.create_task(run())

    async def main():
        ai._lanes = ai._Lanes()
        bgs = [asyncio.create_task(call(f"bg{i}", True, 0.3)) for i in range(5)]
        await asyncio.sleep(0.05)
        t0 = asyncio.get_running_loop().time()
        await call("user", False, 0.01)
        waited = asyncio.get_running_loop().time() - t0
        await asyncio.gather(*bgs)
        return waited

    waited = asyncio.run(main())
    assert waited < 0.1                       # the person did not wait for the background calls
    assert order.index("user") <= 2           # only the two background calls already running came first
    db.set_setting("upgrade_parallel", 8)


def test_practice_rounds_and_fix_question(client):
    """練習任務: each round has its own score; right answers are not asked again. 說明後更新: a question
    is redone with the parent's description and keeps it."""
    from app import db
    exam = next(e for e in client.get("/api/exams").json() if e["status"] == "ready")
    qs = db.q("SELECT id FROM questions WHERE exam_id=? ORDER BY seq LIMIT 3", (exam["id"],))
    ids = [q["id"] for q in qs]
    src = f"exam:{exam['id']}"
    s = client.post("/api/practice/sessions", json={"title": "第一次", "source": src, "ids": ids}).json()["id"]
    client.post(f"/api/practice/sessions/{s}/result", json={"qid": ids[0], "correct": True})
    client.post(f"/api/practice/sessions/{s}/result", json={"qid": ids[1], "correct": False})
    client.post(f"/api/practice/sessions/{s}/result", json={"qid": ids[1], "correct": True})   # only the first try counts
    row = client.post(f"/api/practice/sessions/{s}/finish").json()
    assert row["done"] == 2 and row["correct"] == 1 and row["pct"] == 50 and row["finished"]
    pr = client.get("/api/practice/progress", params={"source": src}).json()
    assert pr["right"] == [ids[0]] and pr["wrong"] == [ids[1]] and pr["rounds"][-1]["pct"] == 50
    assert any(r["id"] == s for r in client.get("/api/practice/sessions").json())
    r = client.post(f"/api/questions/{ids[0]}/fix", json={"note": "第一個循環有 5 個灰色"})
    assert r.status_code == 200 and wait(client, r.json()["job_id"])["status"] == "done"
    q = db.one("SELECT fix_note, upgraded_at FROM questions WHERE id=?", (ids[0],))
    assert "5 個灰色" in q["fix_note"] and q["upgraded_at"]
    assert client.post(f"/api/questions/{ids[0]}/fix", json={"note": ""}).status_code == 400


def test_sketch_for_question_without_figure(client):
    from tests import fake_ai
    q = client.get("/api/questions?limit=1").json()["items"][0]
    r = client.post(f"/api/questions/{q['id']}/sketch", json={}).json()
    assert "<svg" in r["svg"] and r["caption"]
    n = fake_ai.calls.count("sketch")
    assert client.post(f"/api/questions/{q['id']}/sketch", json={}).json()["svg"] == r["svg"]  # kept
    assert fake_ai.calls.count("sketch") == n
    client.post(f"/api/questions/{q['id']}/sketch", json={"again": True})  # 換一張圖
    assert fake_ai.calls.count("sketch") == n + 1


def test_vertical_exam_word_paper(client):
    """A 直書 original (國語) gives a 直書 landscape Word paper with real 注音 phonetic guides."""
    import zipfile
    from app import db, exporters
    e = client.get("/api/exams").json()[0]
    ex = db.one("SELECT options FROM exams WHERE id=?", (e["id"],))
    db.update("exams", e["id"], {"options": {**(ex["options"] or {}), "layout": {"direction": "vertical", "orientation": "landscape", "zhuyin": True}}})
    qid = db.scalar("SELECT id FROM questions WHERE exam_id=? ORDER BY seq LIMIT 1", (e["id"],))
    db.update("questions", qid, {"stem": "天空下起大雨，「{□|ㄔㄨㄤ}」外 12 個人"})
    qids = [r["id"] for r in db.q("SELECT id FROM questions WHERE exam_id=? ORDER BY seq", (e["id"],))]
    tok = exporters.make_spec(e["title"], qids, {"template": "quiz"}, e["id"])
    assert exporters._specs[tok]["config"]["writing_mode"] == "vertical" and exporters._specs[tok]["config"]["keep_original"]
    xml = zipfile.ZipFile(exporters.docx(tok)).read("word/document.xml").decode()
    assert 'w:textDirection w:val="tbRl"' in xml and "<w:ruby>" in xml and "ㄔㄨㄤ" in xml and "１２" in xml


def test_homework_grading(client, sample):
    """批改功課: homework photographed → every answer checked; nothing is a 錯題 until confirmed."""
    with open(sample, "rb") as f:
        r = client.post("/api/exams", files=[("files", ("hw.jpg", f, "image/jpeg"))],
                        data={"title": "", "subject": "", "remove_handwriting": "true", "remove_color": "true", "mode": "grade"})
    assert r.status_code == 200, r.text
    assert wait(client, r.json()["job_id"])["status"] == "done"
    e = client.get(f"/api/exams/{r.json()['exam_id']}").json()
    g = e["options"]["grade"]
    assert len(g) == len(e["questions"]) and e["options"]["grade_summary"]["wrong"] >= 1  # graded before the 詳解
    wrong = [int(k) for k, v in g.items() if v["result"] == "wrong"]
    assert client.post("/api/questions/bulk", json={"ids": wrong, "action": "set", "data": {"is_wrong": 1}}).json()["ok"]
    after = {q["id"]: q for q in client.get(f"/api/exams/{e['id']}").json()["questions"]}
    assert all(after[i]["is_wrong"] for i in wrong)


def test_grade_exam_in_bank(client, sample):
    """批改 an exam already in the bank: with its own answers, and after a redo photographed again."""
    from tests import fake_ai
    e = client.get("/api/exams").json()[0]
    j = wait(client, client.post(f"/api/exams/{e['id']}/grade").json()["job_id"])
    assert j["status"] == "done" and "correct" in j["result"]
    ex = client.get(f"/api/exams/{e['id']}").json()
    assert ex["options"]["grade"] and ex["options"]["grade_source"] == "scan"
    with open(sample, "rb") as f:
        r = client.post(f"/api/exams/{e['id']}/grade/photos", files=[("files", ("redo.jpg", f, "image/jpeg"))])
    assert wait(client, r.json()["job_id"])["status"] == "done"
    ex = client.get(f"/api/exams/{e['id']}").json()
    assert ex["options"]["grade_source"] == "photos" and "read_answers" in fake_ai.calls
    assert all(v["read_as"] in ("③", "8") for v in ex["options"]["grade"].values())
