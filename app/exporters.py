"""Print specs, PDF (headless Chromium renders the same print page the browser
uses, so PDF and on-screen print match) and Word export."""
import asyncio
import json
import re
import secrets
import shutil
import tempfile
import time
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import db
from .config import EXPORT_DIR, MEDIA_DIR, PORT
from .practice import latex_to_text

_specs: dict[str, dict] = {}
_pdf_lock = asyncio.Lock()

DEFAULT_CONFIG = {
    "orientation": "portrait",   # portrait 直印 / landscape 橫印
    "columns": 1,
    "font_size": 14,
    "answer_space": 2,           # blank lines after each question (0-8)
    "answers": "none",           # none / end / inline / separate
    "explanations": False,
    "header": True,              # name / date / score box
    "writing_mode": "horizontal",  # horizontal / vertical (直書)
    "show_figures": True,
    "show_source": False,
    "mode": "questions",         # questions (editable text) / images (cleaned scan)
    "bw": False,
    "group_by_type": False,      # 依題型分大題（一、選擇題 二、填充題 …）
    "numbering": "section",      # section: 每大題重新編號 / continuous: 全卷連續
    "lesson": False,             # 卷首附「觀念教學重點」（先教再練）
    "visuals": True,             # 圖解（列印為靜態圖）
    "template": "",              # 試卷範本 id (app/templates.py); "" = plain layout
    "answer_style": "",          # 解答卷: brief 只有答案 / steps 答案＋詳解 / visual 詳解＋圖解
    "key_only": False,           # only the answer key (解答卷 by itself)
}

ANSWER_STYLE = {"brief": {"explanations": False}, "steps": {"explanations": True, "visuals": False},
                "visual": {"explanations": True, "visuals": True}}


def make_spec(title: str, question_ids: list[int], config: dict, exam_id: int = 0,
              subtitle: str = "", paper_id: int = 0, sample: bool = False) -> str:
    tok = secrets.token_urlsafe(16)
    now = time.time()
    for k in [k for k, v in _specs.items() if now - v["t"] > 6 * 3600]:
        _specs.pop(k, None)
    cfg = {**DEFAULT_CONFIG, **(config or {})}
    cfg.update(ANSWER_STYLE.get(cfg.get("answer_style") or "", {}))
    if exam_id and "keep_original" not in (config or {}):
        # the whole scanned exam: its own sections, numbers and order, like the photographed paper
        n_all = db.scalar("SELECT COUNT(*) FROM questions WHERE exam_id=?", (exam_id,)) or 0
        cfg["keep_original"] = bool(question_ids) and len(question_ids) >= n_all
    if not exam_id and "writing_mode" not in (config or {}) and question_ids:
        # a paper of questions from several exams (錯題卷…): those from 直書 originals print 直書 —
        # the whole paper when they all are, otherwise a 直書 part of their own after the rest
        marks = ",".join("?" * len(question_ids))
        vids = [r["id"] for r in db.q(
            f"SELECT q.id FROM questions q JOIN exams e ON e.id = q.exam_id WHERE q.id IN ({marks}) "
            f"AND json_extract(e.options, '$.layout.direction') = 'vertical'", list(question_ids))]
        if vids and len(set(vids)) == len(set(question_ids)):
            cfg["writing_mode"], cfg["orientation"] = "vertical", "landscape"
        elif vids:
            cfg["vertical_ids"] = vids
    if exam_id and "writing_mode" not in (config or {}):
        # one scanned exam: printed the way the original is (直書 國語 on a landscape page, …)
        e = db.one("SELECT options FROM exams WHERE id=?", (exam_id,))
        lay = ((e or {}).get("options") or {}).get("layout") if isinstance((e or {}).get("options"), dict) else None
        if lay and lay.get("direction") == "vertical":
            cfg["writing_mode"] = "vertical"
            cfg["orientation"] = "landscape" if lay.get("orientation") == "landscape" else "portrait"
    cfg["columns"] = 1  # A4: never two columns
    if cfg.get("key_only") and cfg.get("answers") in ("none", "inline"):
        cfg["answers"] = "separate"
    _specs[tok] = {"t": now, "title": title, "subtitle": subtitle, "question_ids": question_ids,
                   "config": cfg, "exam_id": exam_id, "paper_id": paper_id,
                   "sample": sample}
    return tok


def _figure_sizes(qs: list[dict]):
    """The figures at the size they have on the photographed paper (fig_frac: the figure's share of
    the page width — its crop against the page scan it was cut from), for print.js and the Word file."""
    widths: dict = {}
    for q in qs:
        if not q.get("figure_path") or not q.get("page_id"):
            continue
        if q["page_id"] not in widths:
            widths[q["page_id"]] = db.scalar("SELECT width FROM pages WHERE id=?", (q["page_id"],)) or 0
        f = MEDIA_DIR / q["figure_path"]
        if widths[q["page_id"]] and f.exists():
            try:
                from PIL import Image
                with Image.open(f) as im:
                    q["fig_frac"] = round(min(1.0, im.width / widths[q["page_id"]]), 3)
            except Exception:
                pass


def resolve(tok: str) -> dict | None:
    s = _specs.get(tok)
    if not s:
        return None
    from . import templates
    qs = []
    if s.get("sample"):
        qs = templates.sample_questions()
    for qid in s["question_ids"]:
        r = db.one("SELECT * FROM questions WHERE id=?", (qid,))
        if r:
            qs.append(r)
    cfg = s["config"]
    if cfg.get("keep_original"):
        # as on the original: its headings (a page that continues the last section has none) and numbers
        section = ""
        for i, q in enumerate(qs, 1):
            section = (q.get("section") or "").strip() or section
            q["orig_section"] = q["section"] = section
            q["number"] = str(q.get("number") or "").strip().rstrip(".、") or str(i)
        _figure_sizes(qs)
    elif cfg.get("group_by_type"):
        from .pipeline import number_questions
        nums = number_questions(qs, True, cfg.get("numbering") != "continuous")
        qs.sort(key=lambda q: nums[q["id"]]["order"])
        for q in qs:
            q["orig_section"] = q.get("section") or ""   # the exam's own heading (its points, e.g. 每題 2 分)
            q["section"] = nums[q["id"]]["section"]
            q["number"] = nums[q["id"]]["label"]
    else:
        for i, q in enumerate(qs, 1):
            q["number"] = str(i)
    lesson = None
    if cfg.get("lesson") and s.get("paper_id"):
        lesson = db.scalar("SELECT lesson FROM papers WHERE id=?", (s["paper_id"],)) or None
        if lesson:
            import json as _json
            lesson = _json.loads(lesson)
    pages = []
    if s["exam_id"]:
        pages = db.q("SELECT id, idx, clean_path, scan_path, width, height FROM pages "
                     "WHERE exam_id=? AND status!='error' ORDER BY idx", (s["exam_id"],))
    out = {"title": s["title"], "subtitle": s["subtitle"], "config": s["config"],
           "questions": qs, "pages": pages, "lesson": lesson, "layout": None, "vis_png": s.get("vis_png") or {},
           "fig_png": s.get("fig_png") or {}}
    if cfg.get("template"):
        tpl = templates.get(templates.resolve_id(cfg["template"], s["title"]))
        subj = max((q.get("subject") or "" for q in qs), key=[q.get("subject") for q in qs].count, default="")
        grade = max((q.get("grade") or 0 for q in qs), default=0) or db.scalar("SELECT grade FROM students ORDER BY id LIMIT 1") or ""
        date = db.scalar("SELECT exam_date FROM exams WHERE id=?", (s["exam_id"],)) if s["exam_id"] else ""
        if not date and qs and qs[0].get("exam_id"):
            date = db.scalar("SELECT exam_date FROM exams WHERE id=?", (qs[0]["exam_id"],))
        templates.apply_layout(out, tpl, {"subject": subj, "grade": grade, "date": date or db.today()})
    return out


def chrome_cmd(exe: str, who: str) -> list[str]:
    """Chromium started at low CPU priority (a child's AI question comes first), its profile and
    cache in RAM instead of on the Pi's slow SD card. `who`: one profile per concurrent user."""
    shm = Path("/dev/shm") if Path("/dev/shm").is_dir() else Path(tempfile.gettempdir())
    d = shm / f"ai-review-chrome-{who}"
    return ["ionice", "-c", "3", "nice", "-n", "10", exe, f"--user-data-dir={d}", f"--disk-cache-dir={d}/cache"]


def _chromium() -> str | None:
    for name in ("chromium", "chromium-browser", "google-chrome"):
        p = shutil.which(name)
        if p:
            return p
    return None


async def prepare_visuals(tok: str):
    """For a Word file with 圖解 in its answers: the diagrams as pictures (made once, cached)."""
    from . import vischeck
    s, spec = _specs.get(tok), resolve(tok)
    if not s or not spec or "vis_png" in s:
        return
    cfg = spec["config"]
    # the questions' own figures redrawn as SVG (題目升級) → pictures for Word
    svgq = [q for q in spec["questions"] if q.get("figure_svg") and not q.get("figure_table")]
    if svgq and cfg.get("show_figures", True):
        files = await vischeck.pngs([{"type": "custom_svg", "title": "", "caption": "",
                                      "params_json": json.dumps({"svg": q["figure_svg"]})} for q in svgq])
        s["fig_png"] = {str(q["id"]): str(f) for q, f in zip(svgq, files) if f}
    if not (cfg.get("explanations") and cfg.get("visuals") and cfg.get("answers") != "none"):
        s["vis_png"] = {}
        return
    wanted = [(q["id"], v) for q in spec["questions"] for v in (q.get("visuals") or [])[:2]
              if isinstance(v, dict) and v.get("type") and not v.get("pending")]
    files = await vischeck.pngs([v for _, v in wanted]) if wanted else []
    out: dict = {}
    for (qid, v), f in zip(wanted, files):
        if f:
            out.setdefault(str(qid), []).append((str(f), plain(v.get("caption") or "")))
    s["vis_png"] = out


async def pdf(tok: str) -> Path:
    spec = resolve(tok)
    if spec and (spec.get("layout") or {}).get("custom"):  # the parent's own Word template: PDF from that Word file
        from . import templates
        await prepare_visuals(tok)
        return await templates.to_pdf(await asyncio.to_thread(docx, tok))
    exe = _chromium()
    if not exe:
        raise RuntimeError("找不到 Chromium，無法產生 PDF（可改用瀏覽器的列印 → 另存 PDF）")
    out = EXPORT_DIR / f"{tok}.pdf"
    url = f"http://127.0.0.1:{PORT}/print.html?t={tok}&pdf=1"
    async with _pdf_lock:
        proc = await asyncio.create_subprocess_exec(
            *chrome_cmd(exe, "pdf"), "--headless=new", "--no-sandbox", "--disable-gpu", "--hide-scrollbars",
            "--no-pdf-header-footer", "--run-all-compositor-stages-before-draw",
            "--virtual-time-budget=20000", f"--print-to-pdf={out}", url,
            stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL)
        try:
            await asyncio.wait_for(proc.wait(), timeout=120)
        except asyncio.TimeoutError:
            proc.kill()
            raise RuntimeError("產生 PDF 逾時")
    if not out.exists():
        raise RuntimeError("產生 PDF 失敗")
    return out


# --- Word -------------------------------------------------------------------

_RUBY = re.compile(r"\{([^{}|]+)\|([^{}]+)\}")


def plain(s: str) -> str:
    s = re.sub(r"\\n(?![A-Za-z])", "\n", s or "")  # "\n" written as text by the AI (not \neq, \ne…)
    s = _RUBY.sub(lambda m: f"{m.group(1)}({m.group(2)})", s)
    s = re.sub(r"\$([^$]+)\$", lambda m: latex_to_text(m.group(1)), s)
    return s.replace("**", "")


CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩"


import contextvars
_vertical = contextvars.ContextVar("docx_vertical", default=False)  # building a 直書 Word paper
_FULLWIDTH = str.maketrans("0123456789", "０１２３４５６７８９")


def _add_text(p, text: str, size: int, bold=False, color=None):
    # **bold** segments; {字|注音} becomes a real phonetic guide (Word 注音標示), the 注音 beside
    # the character the way the paper prints it (to the right in 直書, above in 橫書)
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        last = 0
        for m in [*_RUBY.finditer(part), None]:
            seg = part[last:m.start()] if m else part[last:]
            if seg:
                t = plain(seg)
                if _vertical.get():
                    t = t.translate(_FULLWIDTH)
                r = p.add_run(t)
                _font(r, size, bold or (i % 2 == 1), color)
            if m:
                _ruby(p, m.group(1), m.group(2).replace(" ", ""), size, bold or (i % 2 == 1))
                last = m.end()


def _font(r, size: int, bold=False, color=None):
    r.font.size = Pt(size)
    r.bold = bold
    r.font.name = "Microsoft JhengHei"
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "標楷體")
    if color:
        r.font.color.rgb = RGBColor(*color)


def _ruby(p, base: str, rt: str, size: int, bold=False):
    from docx.oxml import OxmlElement

    def el(tag, val=None):
        e = OxmlElement(tag)
        if val is not None:
            e.set(qn("w:val"), str(val))
        return e

    def run(text, pts):
        r = el("w:r")
        rpr = el("w:rPr")
        fonts = el("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            fonts.set(qn(a), "標楷體")
        rpr.append(fonts)
        if bold:
            rpr.append(el("w:b"))
        rpr.append(el("w:sz", int(pts * 2)))
        rpr.append(el("w:szCs", int(pts * 2)))
        r.append(rpr)
        t = el("w:t")
        t.text = text
        r.append(t)
        return r

    small = max(6, round(size * 0.42))
    ruby = el("w:ruby")
    pr = el("w:rubyPr")
    for tag, val in (("w:rubyAlign", "distributeSpace"), ("w:hps", small * 2), ("w:hpsRaise", size * 2),
                     ("w:hpsBaseText", size * 2), ("w:lid", "zh-TW")):
        pr.append(el(tag, val))
    ruby.append(pr)
    rt_el, base_el = el("w:rt"), el("w:rubyBase")
    rt_el.append(run(rt, small))
    base_el.append(run(base, size))
    ruby.append(rt_el)
    ruby.append(base_el)
    holder = p.add_run()
    holder._r.append(ruby)


def docx(tok: str) -> Path:
    from . import templates
    spec = resolve(tok)
    if not spec:
        raise RuntimeError("列印設定已過期，請重新產生")
    cfg = spec["config"]
    lay = spec.get("layout") or {}
    if lay.get("custom"):
        out = EXPORT_DIR / f"{tok}.docx"
        return templates.fill_docx(templates.TPL_DIR / f"{lay['template']}.docx", spec, out, plain)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    if cfg["orientation"] == "landscape":
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.6))
    _vertical.set(cfg.get("writing_mode") == "vertical" and not cfg.get("key_only"))
    if cfg.get("writing_mode") == "vertical" and not cfg.get("key_only"):
        # 直書 like the original: text top to bottom, lines right to left (the answer key stays 橫書)
        from docx.oxml import OxmlElement
        td = OxmlElement("w:textDirection")
        td.set(qn("w:val"), "tbRl")
        grid = sec._sectPr.find(qn("w:docGrid"))  # Word wants it before docGrid
        if grid is not None:
            grid.addprevious(td)
        else:
            sec._sectPr.append(td)
    size = int(cfg["font_size"])  # A4: always one column (the way exam software prints)
    if cfg.get("key_only"):  # 解答卷 by itself: no name / score form
        if lay:
            templates._page_number_footer(sec, size)
    elif lay and cfg.get("writing_mode") == "vertical":
        # 直書: tables turn sideways in a vertical page — a title and a name line, like the original
        templates._page_number_footer(sec, size)
        if cfg["header"]:
            _add_text(doc.add_paragraph(), spec["title"], size + 3, bold=True)
            if spec.get("subtitle"):
                _add_text(doc.add_paragraph(), spec["subtitle"], size + 1)
            _add_text(doc.add_paragraph(), "班級：＿＿＿　座號：＿＿＿　姓名：＿＿＿＿＿　得分：＿＿＿", size)
    elif lay:
        templates._page_number_footer(sec, size)
        style = lay.get("header_style", "simple") if cfg["header"] else "none"
        templates.header_block(doc, {"header_style": style}, size, spec["title"], spec.get("subtitle", ""))
        if lay.get("instructions"):
            _add_text(doc.add_paragraph(), "作答說明：" + lay["instructions"], size - 1)
        if lay.get("score_table"):
            templates.score_table(doc, lay["score_table"], size)
            doc.add_paragraph()
    elif not cfg.get("key_only"):
        t = doc.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_text(t, spec["title"], size + 6, bold=True)
        if spec.get("subtitle"):
            st = doc.add_paragraph()
            st.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_text(st, spec["subtitle"], size)
    if cfg["header"] and not lay and not cfg.get("key_only"):
        h = doc.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_text(h, "座號：＿＿＿　姓名：＿＿＿＿＿＿　　日期：＿＿＿＿＿＿　　得分：＿＿＿＿", size)
    les = spec.get("lesson")
    if les:
        p = doc.add_paragraph()
        _add_text(p, "觀念教學重點（先看懂，再解題）", size + 2, bold=True)
        if les.get("intro"):
            _add_text(doc.add_paragraph(), les["intro"], size)
        for i, c in enumerate(les.get("concepts", []), 1):
            _add_text(doc.add_paragraph(), f"{i}. {c.get('title', '')}", size + 1, bold=True)
            for label, key in (("", "explain"), ("規則：", "rule"), ("例題：", "example_question"),
                               ("解法：", "example_solution"), ("小心：", "mistake")):
                if c.get(key):
                    _add_text(doc.add_paragraph(), label + c[key], size - 1)
        doc.add_page_break()
    from . import docxq
    figs = spec.get("fig_png") or {}
    vset = set(cfg.get("vertical_ids") or [])

    def questions(items):
        section = None
        for n, q in items:
            if q["section"] and q["section"] != section:
                section = q["section"]
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.keep_with_next = True
                _add_text(p, section, size + 1, bold=True)
            choice = docxq.add_question(doc, q, n, cfg, size, _add_text, plain, MEDIA_DIR, figs.get(str(q["id"])))
            if cfg["answers"] == "inline":
                ap = doc.add_paragraph()
                _add_text(ap, f"答案：{q['answer']}", size, color=(200, 30, 30))
                if cfg["explanations"] and q["explanation"]:
                    ep = doc.add_paragraph()
                    _add_text(ep, f"詳解：{q['explanation']}", size - 2, color=(60, 60, 160))
            elif not choice:  # room to write: only where it is needed
                for _ in range(int(cfg["answer_space"])):
                    doc.add_paragraph()

    allq = [] if cfg.get("key_only") else list(enumerate(spec["questions"], 1))
    questions([x for x in allq if x[1]["id"] not in vset])
    vpart = [x for x in allq if x[1]["id"] in vset]
    if vpart:
        # the questions from 直書 originals: their own landscape 直書 section
        from docx.enum.section import WD_SECTION
        from docx.oxml import OxmlElement
        vs = doc.add_section(WD_SECTION.NEW_PAGE)
        vs.orientation = WD_ORIENT.LANDSCAPE
        vs.page_width, vs.page_height = Cm(29.7), Cm(21)
        td = OxmlElement("w:textDirection")
        td.set(qn("w:val"), "tbRl")
        grid = vs._sectPr.find(qn("w:docGrid"))
        grid.addprevious(td) if grid is not None else vs._sectPr.append(td)
        _vertical.set(True)
        _add_text(doc.add_paragraph(), "國語（直書）", size + 2, bold=True)
        questions(vpart)
    if cfg["answers"] in ("end", "separate"):
        if not cfg.get("key_only") and _vertical.get():
            # after 直書 questions the answer key is 橫書 again: its own section (portrait A4)
            from docx.enum.section import WD_SECTION
            ks = doc.add_section(WD_SECTION.NEW_PAGE)
            td = ks._sectPr.find(qn("w:textDirection"))
            if td is not None:
                ks._sectPr.remove(td)
            if vset:
                ks.orientation = WD_ORIENT.PORTRAIT
                ks.page_width, ks.page_height = Cm(21), Cm(29.7)
            _vertical.set(False)
        elif not cfg.get("key_only"):
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_text(p, f"{spec['title']}　{'解答卷' if cfg['answers'] == 'separate' else '參考答案'}", size + 4, bold=True)
        ksec = None
        if not cfg["explanations"]:  # 只有答案: a compact answer grid, five per row — one per section, labelled
            per = 5
            groups: list = []
            for q in spec["questions"]:
                sec = re.sub(r"（.*$", "", q.get("section") or "")
                if not groups or groups[-1][0] != sec:
                    groups.append((sec, []))
                groups[-1][1].append(q)
            for sec, qs in groups:
                if sec:
                    _add_text(doc.add_paragraph(), sec, size, bold=True)
                t = doc.add_table(rows=(len(qs) + per - 1) // per, cols=per)
                t.style = "Table Grid"
                for i, q in enumerate(qs):
                    cell = t.rows[i // per].cells[i % per]
                    cell.text = ""
                    _add_text(cell.paragraphs[0], f"{q.get('number') or i + 1}. {plain(q['answer'] or '')}", size - 1)
        for n, q in enumerate(spec["questions"] if cfg["explanations"] else [], 1):
            if ksec != q.get("section") and q.get("section"):
                ksec = q["section"]
                _add_text(doc.add_paragraph(), re.sub(r"（.*$", "", ksec), size, bold=True)
            ap = doc.add_paragraph()
            _add_text(ap, f"{q.get('number') or n}. {q['answer']}", size, bold=True)
            if cfg["explanations"]:
                for line in templates.explain_text(q).split("\n"):
                    if line.strip():
                        ep = doc.add_paragraph()
                        ep.paragraph_format.left_indent = Cm(0.6)
                        _add_text(ep, line, size - 2, color=(40, 50, 120))
                docxq.option_notes(doc, q, size, _add_text)
                for path, cap in (spec.get("vis_png") or {}).get(str(q["id"]), []):
                    try:
                        doc.add_picture(path, width=Cm(11))
                    except Exception:
                        pass
                    if cap:
                        cp = doc.add_paragraph()
                        _add_text(cp, cap, size - 3, color=(90, 90, 90))
    out = EXPORT_DIR / f"{tok}.docx"
    doc.save(out)
    return out
