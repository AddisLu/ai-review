"""試卷範本: how a scanned exam (or any paper) is laid out again — 小考, 期中考,
模擬考, 私中複習卷 … Each built-in template is a set of print options plus a
header layout (title line, 學校/班級/座號 form, points per section, score table).

A parent can download any template as a Word file, change it in Word (fonts,
school name, logo, header, margins …) and upload it again. The {{…}} marks in
that file are where the questions and exam details go; everything else is kept
exactly as she made it, and the PDF / print of that template is converted from
the same filled-in Word file, so what comes out matches what she designed."""
import asyncio
import copy
import datetime
import re
import secrets
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from . import db
from .config import DATA_DIR, MEDIA_DIR

TPL_DIR = DATA_DIR / "templates"
TPL_DIR.mkdir(parents=True, exist_ok=True)

POINTS = {"是非題": 2, "選擇題": 2, "填充題": 3, "配合題": 2, "連連看": 2, "圈選題": 2, "改錯題": 3,
          "計算題": 4, "應用題": 5, "看圖作答": 4, "作圖題": 4, "簡答題": 4, "造句題": 3, "閱讀測驗": 4, "其他": 3}

BUILTIN = [
    {"id": "quiz", "name": "小考", "icon": "edit",
     "desc": "一張 A4、簡單表頭（座號・姓名・得分），題號連續、答案附在最後。適合隨堂小考、單元小測驗。",
     "config": {"answer_style": "brief", "group_by_type": False, "answer_space": 1, "answers": "end", "explanations": False, "columns": 1,
                "font_size": 14, "header": True, "lesson": False},
     "layout": {"header_style": "simple", "title_format": "{title}"}},
    {"id": "unit", "name": "單元評量", "icon": "book",
     "desc": "依題型分大題（一、選擇題 二、填充題…），每大題標分數，另附解答卷。適合單元複習、平時考。",
     "config": {"answer_style": "steps", "group_by_type": True, "numbering": "section", "answer_space": 2, "answers": "separate",
                "explanations": False, "columns": 1, "font_size": 14, "header": True, "lesson": False},
     "layout": {"header_style": "simple", "title_format": "{title}", "show_points": True}},
    {"id": "midterm", "name": "期中考／期末考", "icon": "doc",
     "desc": "學校段考格式：學年度・學期・年級科目標題，班級／座號／姓名／得分表格、大題得分表，每大題標分數，解答卷另附。",
     "config": {"answer_style": "brief", "group_by_type": True, "numbering": "section", "answer_space": 2, "answers": "separate",
                "explanations": False, "columns": 1, "font_size": 13, "header": True, "lesson": False},
     "layout": {"header_style": "school", "title_format": "{school}{year}學年度第{term}學期　{grade}年級{subject}　{label}",
                "label": "期中評量", "show_points": True, "score_table": True}},
    {"id": "mock", "name": "模擬考", "icon": "bolt",
     "desc": "正式考試感：作答說明（題數、滿分、作答時間）、全卷連續題號、每大題配分，解答卷另附。",
     "config": {"answer_style": "steps", "group_by_type": True, "numbering": "continuous", "answer_space": 2, "answers": "separate",
                "explanations": True, "columns": 1, "font_size": 13, "header": True, "lesson": False},
     "layout": {"header_style": "school", "title_format": "{grade}年級{subject}　模擬考", "show_points": True,
                "minutes": 40, "instructions": "本卷共 {n} 題，滿分 {total} 分，作答時間 {minutes} 分鐘。請把算式和答案寫在題目下方。"}},
    {"id": "private", "name": "私中模擬考（補習班）", "icon": "target",
     "desc": "補習班的私中入學模擬試題：A4 單欄、（　）題號、選項依長度排成一列或兩列、右圖放題目旁、表格重建成表格，全卷連續題號，解答附詳解、選項解析與圖解。",
     "config": {"answer_style": "visual", "group_by_type": True, "numbering": "continuous", "answer_space": 2, "answers": "end",
                "explanations": True, "columns": 1, "font_size": 13, "header": True, "show_source": False, "lesson": False},
     "layout": {"header_style": "simple", "title_format": "{title}"}},
    {"id": "chinese", "name": "國文直書（直印）", "icon": "book",
     "desc": "國語／國文考卷：直書由右而左、A4 直式列印；注音在字的右邊、調號在旁，看注音寫國字用方格；表頭班級・座號・姓名・得分，照原卷的大題和題號，解答卷另附。",
     "config": {"answer_style": "brief", "group_by_type": False, "numbering": "section", "answer_space": 1, "answers": "separate",
                "explanations": False, "columns": 1, "font_size": 15, "header": True, "lesson": False,
                "writing_mode": "vertical", "orientation": "portrait"},
     "layout": {"header_style": "school", "title_format": "{title}"}},
    {"id": "wrong", "name": "錯題複習卷", "icon": "x",
     "desc": "先教再練：卷首「觀念教學重點」與圖解，題目依題型分大題，最後附答案與詳解。",
     "config": {"answer_style": "visual", "group_by_type": True, "numbering": "section", "answer_space": 2, "answers": "end",
                "explanations": True, "columns": 1, "font_size": 14, "header": True, "lesson": True, "visuals": True},
     "layout": {"header_style": "simple", "title_format": "{title}"}},
]

# 解答卷: only the answers / answers with worked steps / steps plus the solving diagrams (SVG)
ANSWER_STYLES = [("brief", "只有答案", "答案表格，一眼對答案"), ("steps", "答案＋詳解", "每題的解題步驟"),
                 ("visual", "詳解＋圖解", "步驟加上解題圖解（線段圖、分數條…）")]

# the marks a Word template may use (shown on the template page as help)
FIELDS = {"{{試卷名稱}}": "試卷標題", "{{副標題}}": "原本的試卷名稱", "{{科目}}": "科目", "{{年級}}": "年級",
          "{{日期}}": "考試日期", "{{學年度}}": "例如 115", "{{學期}}": "一 或 二", "{{學校}}": "學校名稱",
          "{{題數}}": "題目數", "{{總分}}": "總分", "{{考試時間}}": "作答分鐘數"}
BLOCK = {"{{大題}}": "大題標題（一、選擇題…）", "{{題號}}": "題號", "{{題目}}": "題目文字", "{{選項}}": "選擇題選項",
         "{{作答區}}": "作答空白行", "{{大題得分表}}": "大題得分表", "{{解答標題}}": "解答卷標題",
         "{{答案}}": "答案（解答卷）", "{{詳解}}": "詳解（解答卷）", "{{圖解}}": "解題圖解（解答卷）"}


# --- the list ----------------------------------------------------------------------------------------

def _custom() -> list[dict]:
    return [t for t in (db.get_setting("templates_custom") or []) if (TPL_DIR / f"{t['id']}.docx").exists()]


# 整理清單: the parent's own exam types (e.g. 補習班週考) built on a template, with keywords that
# tell 「AI 自動判斷」 which exams belong to it; types she does not use can be hidden from the menus
TYPE_KEYS = {"answer_style", "font_size", "numbering", "group_by_type", "answer_space", "answers"}  # A4: always one column


def _types() -> list[dict]:
    return db.get_setting("template_types") or []


def hidden() -> set:
    return set(db.get_setting("template_hidden") or [])


def all_templates() -> list[dict]:
    out = [{**t, "kind": "builtin"} for t in BUILTIN]
    for t in _types():
        base = get(t.get("base") or "quiz", builtin_only=True)
        out.append({**base, "id": t["id"], "name": t["name"], "kind": "type", "base": base["id"], "icon": base.get("icon"),
                    "keywords": t.get("keywords", []),
                    "config": {**base["config"], **{k: v for k, v in (t.get("config") or {}).items() if k in TYPE_KEYS}},
                    "desc": t.get("desc") or f"自己設定的類型（以「{base['name']}」為基礎）"})
    for t in _custom():
        base = get(t.get("base") or "quiz")
        out.append({**base, **t, "kind": "custom", "config": base["config"], "layout": {**base["layout"], **t.get("layout", {})},
                    "desc": t.get("desc") or f"從「{base['name']}」修改的 Word 範本"})
    hid = hidden()
    for t in out:
        t["hidden"] = t["id"] in hid
    return out


TYPE_WORDS = ["段考", "期中", "期末", "週考", "周考", "月考", "小考", "隨堂", "模擬", "會考", "私中", "入學", "複習", "評量", "練習卷", "作業"]


def save_type(body: dict) -> dict:
    """A type from just a name works: the layout is picked from the name, its exam words become
    keywords, and the AI matches scanned exams to it by name too."""
    name = str(body.get("name") or "").strip()[:20]
    if not name:
        raise ValueError("請輸入類型名稱")
    base = body.get("base")
    if not any(t["id"] == base for t in BUILTIN):
        base = guess(name) if guess(name) in {t["id"] for t in BUILTIN} else "quiz"
        if base == "wrong":
            base = "quiz"
    kws = [k.strip()[:12] for k in (body.get("keywords") or []) if str(k).strip()][:8]
    if "keywords" not in body:
        kws = [w for w in TYPE_WORDS if w in name and w not in ("複習", "作業")] or [name]
    cfg = {k: v for k, v in (body.get("config") or {}).items() if k in TYPE_KEYS}
    types = _types()
    tid = body.get("id")
    cur = next((t for t in types if t["id"] == tid), None) if tid else None
    if cur:
        cur.update({"name": name, "base": base, "keywords": kws, "config": cfg})
    else:
        cur = {"id": "t" + secrets.token_hex(4), "name": name, "base": base, "keywords": kws, "config": cfg, "created_at": db.now()}
        types.append(cur)
    db.set_setting("template_types", types)
    return cur


def delete_type(tid: str):
    db.set_setting("template_types", [t for t in _types() if t["id"] != tid])
    db.set_setting("template_for", {k: v for k, v in (db.get_setting("template_for") or {}).items() if k != tid and v != tid})


def set_hidden(tid: str, on: bool):
    h = hidden()
    (h.add if on else h.discard)(tid)
    db.set_setting("template_hidden", sorted(h))


def get(tid: str | None, builtin_only: bool = False) -> dict:
    for t in BUILTIN:
        if t["id"] == tid:
            return t
    if not builtin_only:
        for t in all_templates():
            if t["id"] == tid:
                return t
    return BUILTIN[0]


def guess(title: str, subject: str = "") -> str:
    """AI 自動判斷: pick the template from the exam's name (her own types' keywords first)."""
    t = title or ""
    for ty in _types():
        if any(k and k in t for k in ty.get("keywords", [])):
            return ty["id"]
    if re.search(r"錯題|訂正", t):
        return "wrong"
    if re.search(r"私中|入學|資優", t):
        return "private"
    if re.search(r"模擬|會考", t):
        return "mock"
    if re.search(r"期中|期末|段考|定期", t):
        return "midterm"
    if re.search(r"小考|隨堂|小測", t):
        return "quiz"
    if re.search(r"單元|評量|平時考|複習", t):
        return "unit"
    return "quiz"


def exam_category(exam: dict) -> str:
    """Which kind of exam this is: the template chosen when it was scanned, else from its name
    (her own types' keywords, then the AI's pick among her types, then the built-in rules)."""
    opts = exam.get("options") or {}
    if isinstance(opts, str):
        import json as _json
        opts = _json.loads(opts or "{}")
    title = exam.get("title") or ""
    kw = next((ty["id"] for ty in _types() if any(k and k in title for k in ty.get("keywords", []))), None)
    ai_kind = opts.get("kind_ai")
    if (not opts.get("template") or opts.get("template") == "auto") and not kw and ai_kind and any(t["id"] == ai_kind for t in _types()):
        return ai_kind
    tid = opts.get("template")
    if tid and tid != "auto" and any(t["id"] == tid for t in BUILTIN + _types()):
        return tid
    if tid and tid != "auto":  # an uploaded Word template: the kind it was made from
        c = next((t for t in _custom() if t["id"] == tid), None)
        if c:
            return c.get("base") or "quiz"
    return guess(exam.get("title") or "")


def examples(student_id: int) -> dict:
    """Real exams of each kind, for previews (a kind without any has none — nothing is made up).
    The viewed child's exams first, then the rest of the family's."""
    out: dict = {}
    for e in db.q("SELECT id, title, options, (SELECT COUNT(*) FROM questions q WHERE q.exam_id=e.id) AS n FROM exams e "
                  "WHERE status='ready' ORDER BY (student_id=?) DESC, id DESC", (student_id,)):
        if e["n"]:
            out.setdefault(exam_category(e), []).append({"id": e["id"], "title": e["title"], "n": e["n"]})
    return out


def mapping() -> dict:
    """這類考卷用哪個範本: {kind: template id}, e.g. every 期中考 in the parent's own Word template."""
    ids = {t["id"] for t in all_templates()}
    return {k: v for k, v in (db.get_setting("template_for") or {}).items() if v in ids}


def for_kind(kind: str) -> str:
    return mapping().get(kind) or kind


def resolve_id(tid: str | None, title: str = "") -> str:
    if not tid or tid == "auto":
        return for_kind(guess(title))
    return tid if any(t["id"] == tid for t in all_templates()) else for_kind(guess(title))


def for_exam(exam: dict) -> dict:
    """What this scanned exam turns into, and why (shown on the exam page)."""
    opts = exam.get("options") or {}
    chosen = opts.get("template") if isinstance(opts, dict) else None
    kind = exam_category(exam)
    explicit = bool(chosen and chosen != "auto" and any(t["id"] == chosen for t in all_templates()))
    tid = chosen if explicit else for_kind(kind)
    t = get(tid)
    return {"id": t["id"], "name": t["name"], "kind": kind, "kind_name": get(kind)["name"],
            "custom": t.get("kind") == "custom", "auto": not explicit, "config": t["config"]}


def save_custom(data: bytes, name: str, base: str) -> dict:
    """An uploaded Word template: checked, stored, listed."""
    tmp = Path(tempfile.mkdtemp()) / "t.docx"
    tmp.write_bytes(data)
    try:
        doc = Document(str(tmp))
    except Exception:
        raise ValueError("這不是 Word（.docx）檔，請用 Word 另存成 .docx 再上傳")
    text = _all_text(doc)
    if "{{題目}}" not in text:
        raise ValueError("範本裡找不到 {{題目}} 標記（AI 要在那裡放題目）。請從「下載 Word 範本」開始修改，保留 {{…}} 標記。")
    cat = (doc.core_properties.category or "")
    if cat.startswith("ai-review:") and not base:
        base = cat.split(":", 1)[1]
    base = base if any(t["id"] == base for t in BUILTIN) else "quiz"
    tid = "c" + secrets.token_hex(5)
    shutil.move(str(tmp), TPL_DIR / f"{tid}.docx")
    meta = {"id": tid, "name": (name or "我的範本").strip()[:30], "base": base, "created_at": db.now(),
            "marks": sorted(m for m in list(FIELDS) + list(BLOCK) if m in text)}
    db.set_setting("templates_custom", (db.get_setting("templates_custom") or []) + [meta])
    return meta


def delete_custom(tid: str):
    db.set_setting("templates_custom", [t for t in (db.get_setting("templates_custom") or []) if t["id"] != tid])
    (TPL_DIR / f"{tid}.docx").unlink(missing_ok=True)


def _all_text(doc) -> str:
    parts = [p.text for p in _paragraphs(doc)]
    return "\n".join(parts)


# --- layout values (title, points, score table) -----------------------------------------------------

def school_year(date: str) -> tuple[int, str]:
    try:
        d = datetime.date.fromisoformat((date or "")[:10])
    except ValueError:
        d = datetime.date.today()
    year = d.year - 1911 - (0 if d.month >= 8 else 1)
    return year, "一" if d.month >= 8 or d.month == 1 else "二"


def apply_layout(spec: dict, tpl: dict, info: dict) -> dict:
    """Fill in the template's title / points / header for this paper (in place)."""
    lay = dict(tpl.get("layout") or {})
    qs = spec["questions"]
    pts = {**POINTS, **(lay.get("points") or {})}
    total = 0
    if lay.get("show_points"):
        for q in qs:
            # the points printed on the original paper (每題 10 分) win; defaults only when none were read
            m = re.search(r"每[題格小]\s*(\d+(?:\.\d+)?)\s*分", (q.get("orig_section") or "") + (q.get("section") or ""))
            q["points"] = (float(m.group(1)) if "." in m.group(1) else int(m.group(1))) if m else pts.get(q.get("qtype") or "其他", 3)
            total += q["points"]
        per: dict = {}
        for q in qs:
            if q.get("section"):
                per.setdefault(q["section"], []).append(q["points"])
        names = {}
        for sec, ps in per.items():
            if re.search(r"\d+\s*分", sec):  # the original heading already says its points
                names[sec] = sec
                continue
            each = f"每題 {ps[0]} 分，" if len(set(ps)) == 1 else ""
            names[sec] = re.sub(r"（共 \d+ 題）$", "", sec) + f"（{each}共 {sum(ps)} 分）"
        for q in qs:
            if q.get("section"):
                q["section"] = names[q["section"]]
    year, term = school_year(info.get("date", ""))
    vals = {"title": spec["title"], "subject": info.get("subject", ""), "grade": info.get("grade", ""),
            "year": year, "term": term, "school": lay.get("school", ""),
            "label": "期末評量" if "期末" in spec["title"] else "期中評量" if "期中" in spec["title"] else lay.get("label", ""),
            "n": len(qs), "total": total, "minutes": lay.get("minutes", 40)}
    fmt = lambda s: re.sub(r"\s{2,}", "　", str(s or "").format(**vals)).strip("　 ")
    original = spec["title"]
    spec["title"] = fmt(lay.get("title_format") or "{title}") or original
    if lay.get("subtitle_format"):
        spec["subtitle"] = spec.get("subtitle") or fmt(lay["subtitle_format"])
    elif "{title}" not in (lay.get("title_format") or "{title}") and not spec.get("subtitle"):
        spec["subtitle"] = original
    if spec["subtitle"] == spec["title"]:
        spec["subtitle"] = ""
    sections = list(dict.fromkeys(q["section"] for q in qs if q.get("section")))
    spec["layout"] = {
        "template": tpl["id"], "name": tpl["name"], "custom": tpl.get("kind") == "custom",
        "header_style": lay.get("header_style", "simple"),
        "instructions": fmt(lay.get("instructions", "")),
        "score_table": [re.match(r"^(\S+?)、", s).group(1) if re.match(r"^(\S+?)、", s) else str(i + 1)
                        for i, s in enumerate(sections)] if lay.get("score_table") and sections else [],
        "total": total,
        "vars": {"{{試卷名稱}}": spec["title"], "{{副標題}}": spec.get("subtitle", ""), "{{科目}}": vals["subject"],
                 "{{年級}}": f"{vals['grade']}" if vals["grade"] else "", "{{日期}}": info.get("date", ""),
                 "{{學年度}}": str(year), "{{學期}}": term, "{{學校}}": vals["school"], "{{題數}}": str(len(qs)),
                 "{{總分}}": str(total) if total else "", "{{考試時間}}": str(vals["minutes"])},
    }
    return spec


# --- Word: built-in template documents ----------------------------------------------------------------

def _font(run, size, bold=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = "Microsoft JhengHei"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "標楷體")
    if color:
        run.font.color.rgb = RGBColor(*color)


def _para(doc, text="", size=14, bold=False, align=None, color=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if text:
        _font(p.add_run(text), size, bold, color)
    return p


def _cell(cell, text, size, bold=False):
    cell.text = ""
    _font(cell.paragraphs[0].add_run(text), size, bold)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _page_number_footer(sec, size):
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run("第 "), size - 3)
    r = p.add_run()
    for tag, attr in (("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", None), ("w:fldChar", {"w:fldCharType": "end"})):
        el = r._element.makeelement(qn(tag), {qn(k): v for k, v in (attr or {}).items()})
        if tag == "w:instrText":
            el.text = "PAGE"
        r._element.append(el)
    _font(r, size - 3)
    _font(p.add_run(" 頁"), size - 3)


def header_block(doc, layout: dict, size: int, title: str, subtitle: str):
    """Title, subtitle and the name/score form of a built-in template."""
    _para(doc, title, size + 6, True, WD_ALIGN_PARAGRAPH.CENTER)
    if subtitle:
        _para(doc, subtitle, size, False, WD_ALIGN_PARAGRAPH.CENTER)
    style = layout.get("header_style", "simple")
    if style == "school":
        t = doc.add_table(rows=1, cols=8)
        t.style = "Table Grid"
        for i, (k, w) in enumerate((("班級", 1.6), ("", 2.2), ("座號", 1.6), ("", 1.6), ("姓名", 1.6), ("", 4.2), ("得分", 1.6), ("", 3.0))):
            _cell(t.rows[0].cells[i], k, size - 1, bool(k))
            t.rows[0].cells[i].width = Cm(w)
        _para(doc, "", 6)
    elif style == "simple":
        _para(doc, "座號：＿＿＿　姓名：＿＿＿＿＿＿＿　日期：＿＿＿＿＿　得分：＿＿＿＿", size, False, WD_ALIGN_PARAGRAPH.CENTER)
    return style


def score_table(doc, secs: list, size: int, anchor=None):
    if not secs:
        return None
    t = doc.add_table(rows=2, cols=len(secs) + 2)
    t.style = "Table Grid"
    for i, s in enumerate(["大題"] + secs + ["總分"]):
        _cell(t.rows[0].cells[i], s, size - 2, True)
    _cell(t.rows[1].cells[0], "得分", size - 2, True)
    if anchor is not None:
        anchor.addprevious(t._tbl)
    return t


def template_docx(tid: str, out: Path) -> Path:
    """The editable Word version of a template: the same look as its print, with
    {{…}} marks where the AI puts each exam's details and questions."""
    tpl = get(tid)
    if tpl.get("kind") == "custom":
        shutil.copy(TPL_DIR / f"{tid}.docx", out)
        return out
    cfg, lay = tpl["config"], tpl["layout"]
    size = int(cfg.get("font_size", 14))
    doc = Document()
    doc.core_properties.category = f"ai-review:{tid}"
    doc.core_properties.title = f"{tpl['name']} 範本"
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.6))
    _page_number_footer(sec, size)
    _para(doc, "【範本說明】{{…}} 是 AI 會自動填入的地方：可以改字型、大小、顏色、位置，也可以加學校 Logo、頁首頁尾、表格；"
               "只要保留 {{題目}}（和需要的其他標記）就好。這段說明以【範本說明】開頭，輸出時會自動刪掉。",
          9, color=(120, 120, 120))
    title = "{{試卷名稱}}"
    style = header_block(doc, lay, size, title, "{{副標題}}")
    if lay.get("instructions"):
        _para(doc, "作答說明：本卷共 {{題數}} 題，滿分 {{總分}} 分，作答時間 {{考試時間}} 分鐘。", size - 1)
    if lay.get("score_table"):
        _para(doc, "{{大題得分表}}", size - 2)
    if cfg.get("group_by_type"):
        _para(doc, "{{大題}}", size + 1, True)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    _font(p.add_run("{{題號}}. "), size, True)
    _font(p.add_run("{{題目}}"), size)
    op = _para(doc, "{{選項}}", size)
    op.paragraph_format.left_indent = Cm(0.8)
    _para(doc, "{{作答區}}", size)
    a = _para(doc, "{{解答標題}}", size + 4, True, WD_ALIGN_PARAGRAPH.CENTER)
    a.paragraph_format.page_break_before = True
    _para(doc, "{{答案}}", size, True)
    _para(doc, "{{詳解}}", size - 2, color=(60, 60, 160))
    _para(doc, "{{圖解}}", size - 2)
    doc.save(out)
    return out


# --- Word: filling an uploaded template ----------------------------------------------------------------

def _paragraphs(doc):
    """Every paragraph: body, tables (nested), headers and footers."""
    def walk(container):
        for p in container.paragraphs:
            yield p
        for t in getattr(container, "tables", []):
            for row in t.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(doc)
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            try:
                yield from walk(part)
            except Exception:
                pass


def _merge_runs(p, marks):
    """Word splits text into runs anywhere ("{{題" + "目}}"): join the runs a mark spans."""
    runs = p.runs
    if len(runs) < 2:
        return
    text = "".join(r.text for r in runs)
    for m in marks:
        start = text.find(m)
        while start >= 0:
            end = start + len(m)
            pos, idx = 0, []
            for i, r in enumerate(runs):
                a, b = pos, pos + len(r.text)
                if b > start and a < end:
                    idx.append(i)
                pos = b
            if len(idx) > 1:
                first = runs[idx[0]]
                first.text = "".join(runs[i].text for i in idx)
                for i in idx[1:]:
                    runs[i].text = ""
            start = text.find(m, end)


def _replace(p, mapping: dict):
    if not any(m in p.text for m in mapping):
        return
    _merge_runs(p, list(mapping))
    for r in p.runs:
        t = r.text
        for k, v in mapping.items():
            if k in t:
                t = t.replace(k, v)
        if t != r.text:
            r.text = t


def _clone(proto_el, anchor_el, mapping: dict):
    el = copy.deepcopy(proto_el)
    anchor_el.addprevious(el)
    p = Paragraph(el, None)
    _replace(p, mapping)
    return p


def explain_text(q: dict) -> str:
    """詳解 text: the worked steps when there are some, else the explanation."""
    steps = [x for x in (q.get("steps") or []) if x]
    if steps:
        return "\n".join(f"{i}. {x}" for i, x in enumerate(steps, 1))
    return q.get("explanation") or ""


def fill_docx(src: Path, spec: dict, out: Path, plain) -> Path:
    """Put this paper's questions into an uploaded Word template."""
    doc = Document(str(src))
    cfg, lay = spec["config"], spec.get("layout") or {}
    for p in list(_paragraphs(doc)):
        if p.text.strip().startswith("【範本說明】"):
            p._element.getparent().remove(p._element)
    for p in _paragraphs(doc):
        _replace(p, lay.get("vars") or {})
    body = [Paragraph(el, doc._body) for el in doc.element.body.iterchildren() if el.tag == qn("w:p")]

    def proto(mark):
        return next((p for p in body if mark in p.text), None)

    sec_p, q_p, opt_p, space_p = proto("{{大題}}"), proto("{{題目}}"), proto("{{選項}}"), proto("{{作答區}}")
    head_p, ans_p, exp_p, tbl_p = proto("{{解答標題}}"), proto("{{答案}}"), proto("{{詳解}}"), proto("{{大題得分表}}")
    vis_p = proto("{{圖解}}")
    vis_png = spec.get("vis_png") or {}

    def pictures(q, anchor):
        for path, cap in vis_png.get(str(q["id"]), []):
            src = vis_p if vis_p is not None else exp_p if exp_p is not None else q_p
            fp = _clone(src._element, anchor, {})
            for r in fp.runs:
                r.text = ""
            try:
                fp.add_run().add_picture(str(path), width=Cm(11))
            except Exception:
                pass
            if cap and exp_p is not None:
                _clone(exp_p._element, anchor, {"{{詳解}}": cap})
    size = int(cfg.get("font_size", 14))
    if tbl_p is not None:
        score_table(doc, lay.get("score_table") or [], size, tbl_p._element)
        tbl_p._element.getparent().remove(tbl_p._element)
    qprotos = [p for p in (sec_p, q_p, opt_p, space_p) if p is not None]
    anchor = min(qprotos, key=lambda p: list(doc.element.body).index(p._element))._element
    section = None
    for n, q in enumerate([] if cfg.get("key_only") else spec["questions"], 1):
        num = str(q.get("number") or n)
        if q.get("section") and q["section"] != section:
            section = q["section"]
            if sec_p is not None:
                _clone(sec_p._element, anchor, {"{{大題}}": plain(section)})
        lines = plain(q.get("stem") or "").split("\n")
        if q.get("points") and not q.get("section"):
            lines[0] = f"（{q['points']} 分）{lines[0]}"
        for li, line in enumerate(lines):
            _clone(q_p._element, anchor, {"{{題號}}": num, "{{題目}}": line} if li == 0 else
                   {"{{題號}}. ": "　", "{{題號}}": "", "{{題目}}": line})
        figpng = (spec.get("fig_png") or {}).get(str(q["id"]))
        if cfg.get("show_figures", True) and (q.get("figure_table") or {}).get("rows"):
            from . import docxq
            fp = _clone(q_p._element, anchor, {})
            for r in fp.runs:
                r.text = ""
            tb = docxq.figure_table(doc, q["figure_table"], int(cfg.get("font_size", 14)),
                                    lambda p, t, sz, bold=False, color=None: _font(p.add_run(t), sz, bold, color), plain)
            if tb is not None:
                anchor.addprevious(tb._tbl)
            fp._element.getparent().remove(fp._element)
        elif cfg.get("show_figures", True) and (figpng or (q.get("figure_path") and (MEDIA_DIR / q["figure_path"]).exists())):
            fp = _clone(q_p._element, anchor, {})
            for r in fp.runs:
                r.text = ""
            try:
                fp.add_run().add_picture(figpng or str(MEDIA_DIR / q["figure_path"]), width=Cm(8))
            except Exception:
                pass
        for i, o in enumerate(q.get("options") or []):
            mark = "①②③④⑤⑥⑦⑧⑨⑩"[i] if i < 10 else str(i + 1)
            if opt_p is not None:
                _clone(opt_p._element, anchor, {"{{選項}}": f"{mark} {plain(o)}"})
            else:
                _clone(q_p._element, anchor, {"{{題號}}. ": "　", "{{題號}}": "", "{{題目}}": f"{mark} {plain(o)}"})
        if cfg.get("answers") == "inline":
            src_p = ans_p if ans_p is not None else q_p
            _clone(src_p._element, anchor, {"{{答案}}": f"答：{plain(q.get('answer') or '')}", "{{題號}}. ": "",
                                            "{{題號}}": "", "{{題目}}": f"答：{plain(q.get('answer') or '')}"})
            if cfg.get("explanations") and exp_p is not None and explain_text(q):
                _clone(exp_p._element, anchor, {"{{詳解}}": plain(explain_text(q))})
            if cfg.get("explanations"):
                pictures(q, anchor)
        else:
            for _ in range(int(cfg.get("answer_space") or 0)):
                if space_p is not None:
                    _clone(space_p._element, anchor, {"{{作答區}}": ""})
                else:
                    anchor.addprevious(doc.add_paragraph()._element)
    for p in qprotos:
        p._element.getparent().remove(p._element)
    # answer key
    aprotos = [p for p in (head_p, ans_p, exp_p, vis_p) if p is not None]
    if cfg.get("answers") in ("end", "separate") and ans_p is not None:
        anchor = min(aprotos, key=lambda p: list(doc.element.body).index(p._element))._element
        if head_p is not None:
            _clone(head_p._element, anchor, {"{{解答標題}}": f"{spec['title']}　{'解答卷' if cfg.get('answers') == 'separate' else '參考答案'}"})
        for n, q in enumerate(spec["questions"], 1):
            _clone(ans_p._element, anchor, {"{{答案}}": f"{q.get('number') or n}. {plain(q.get('answer') or '')}"})
            if cfg.get("explanations") and exp_p is not None and explain_text(q):
                _clone(exp_p._element, anchor, {"{{詳解}}": plain(explain_text(q))})
            if cfg.get("explanations"):
                pictures(q, anchor)
    elif cfg.get("answers") in ("end", "separate"):
        h = doc.add_paragraph()
        h.paragraph_format.page_break_before = True
        _font(h.add_run(f"{spec['title']}　參考答案"), size + 4, True)
        for n, q in enumerate(spec["questions"], 1):
            _font(doc.add_paragraph().add_run(f"{q.get('number') or n}. {plain(q.get('answer') or '')}"), size)
    for p in aprotos:
        p._element.getparent().remove(p._element)
    for p in _paragraphs(doc):  # any mark left over (e.g. in a header) is emptied
        _replace(p, {m: "" for m in list(FIELDS) + list(BLOCK)})
    doc.save(out)
    return out


# --- Word → PDF → page pictures (for uploaded templates) ------------------------------------------------

_office_lock = asyncio.Lock()


async def to_pdf(docx_path: Path) -> Path:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("這台主機沒有安裝 LibreOffice，無法把 Word 範本轉成 PDF")
    outdir = docx_path.parent
    # its profile in RAM, at low CPU priority (like Chromium: people come first)
    profile = (Path("/dev/shm") if Path("/dev/shm").is_dir() else Path(tempfile.gettempdir())) / "aireview-lo"
    async with _office_lock:
        proc = await asyncio.create_subprocess_exec(
            "nice", "-n", "10", exe, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore", "--convert-to", "pdf",
            "--outdir", str(outdir), str(docx_path), stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL)
        try:
            await asyncio.wait_for(proc.wait(), timeout=120)
        except asyncio.TimeoutError:
            proc.kill()
            raise RuntimeError("Word 轉 PDF 逾時")
    pdf = docx_path.with_suffix(".pdf")
    if not pdf.exists():
        raise RuntimeError("Word 轉 PDF 失敗")
    return pdf


async def pdf_pages(pdf: Path, dpi: int = 110) -> list[Path]:
    prefix = pdf.with_name(pdf.stem + f"_p{dpi}")
    for old in pdf.parent.glob(prefix.name + "-*.png"):
        old.unlink()
    proc = await asyncio.create_subprocess_exec("ionice", "-c", "3", "nice", "-n", "10", "pdftoppm", "-r", str(dpi), "-png", str(pdf), str(prefix),
                                                stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL)
    await asyncio.wait_for(proc.wait(), timeout=120)
    return sorted(pdf.parent.glob(prefix.name + "-*.png"), key=lambda p: int(re.search(r"-(\d+)\.png$", p.name).group(1)))


# --- sample questions (preview before anything is scanned) ---------------------------------------------

SAMPLE = [
    {"qtype": "是非題", "stem": "一個數乘以 $\\frac{3}{4}$，積一定比原來的數小。", "answer": "×",
     "explanation": "如果原來的數是 0，積還是 0，沒有變小。"},
    {"qtype": "選擇題", "stem": "下面哪一個分數和 $\\frac{2}{3}$ 一樣大？", "options": ["$\\frac{4}{9}$", "$\\frac{6}{9}$", "$\\frac{3}{2}$", "$\\frac{2}{6}$"],
     "answer": "②", "explanation": "分子分母同乘 3：$\\frac{2}{3}=\\frac{6}{9}$。"},
    {"qtype": "選擇題", "stem": "一個長方形長 12 公分、寬 8 公分，它的周長是多少公分？", "options": ["20", "40", "96", "48"],
     "answer": "②", "explanation": "(12+8)×2=40（公分）。"},
    {"qtype": "填充題", "stem": "$1.25$ 公升 = （　　）毫升", "answer": "1250", "explanation": "1 公升 = 1000 毫升。"},
    {"qtype": "計算題", "stem": "$3\\frac{1}{2} \\div \\frac{7}{8}$ = ？", "answer": "4",
     "explanation": "$\\frac{7}{2}\\times\\frac{8}{7}=4$。"},
    {"qtype": "應用題", "stem": "媽媽有 120 元，買文具用了全部的 $\\frac{2}{5}$，剩下多少元？", "answer": "72 元",
     "explanation": "剩下 $1-\\frac{2}{5}=\\frac{3}{5}$，$120\\times\\frac{3}{5}=72$（元）。"},
]


_SAMPLE_EXTRA = {
    2: {"steps": ["分子分母同乘 3", "$\\frac{2}{3}=\\frac{2×3}{3×3}=\\frac{6}{9}$", "答案是 ②"],
        "visuals": [{"type": "fraction_bar", "title": "一樣長的分數條", "caption": "2/3 和 6/9 塗色的長度一樣，所以一樣大",
                     "params_json": '{"bars":[{"label":"2/3","parts":3,"shaded":2,"color":"blue"},{"label":"6/9","parts":9,"shaded":6,"color":"orange"}]}'}]},
    3: {"steps": ["周長 = (長 + 寬) × 2", "(12 + 8) × 2 = 40", "答：40 公分"],
        "visuals": [{"type": "area_grid", "title": "長方形繞一圈", "caption": "繞一圈 = 2 個長 + 2 個寬",
                     "params_json": '{"cols":12,"rows":8,"label_w":"長 12 公分","label_h":"寬 8 公分","note":"(12+8)×2 = 40"}'}]},
    5: {"steps": ["帶分數化成假分數：$3\\frac{1}{2}=\\frac{7}{2}$", "除以分數 = 乘以倒數：$\\frac{7}{2}×\\frac{8}{7}$", "約分得 4"],
        "visuals": [{"type": "steps", "title": "分數除法三步驟", "caption": "先化假分數，再乘倒數，最後約分",
                     "params_json": '{"steps":["3又1/2 → 7/2","÷ 7/8 變成 × 8/7","7/2 × 8/7 = 4"]}'}]},
    6: {"steps": ["用掉 2/5，剩下 1 − 2/5 = 3/5", "120 × 3/5 = 72", "答：剩下 72 元"],
        "visuals": [{"type": "bar_model", "title": "把 120 元分成 5 份", "caption": "用掉 2 份、剩下 3 份，1 份是 24 元",
                     "params_json": '{"bars":[{"label":"全部","note":"共 120 元","segments":[{"label":"用掉 2 份","size":2,"highlight":false},{"label":"剩下 3 份 = ?","size":3,"highlight":true}]}]}'}]},
}


def sample_questions() -> list[dict]:
    out = []
    for i, q in enumerate(SAMPLE, 1):
        out.append({"id": -i, "subject": "數學", "grade": 6, "unit": "範例", "section": "", "number": str(i),
                    "options": [], "figure_path": "", "steps": [], "visuals": [], "bar_model": None, **q,
                    **{k: ([{**v, "checked": True} for v in val] if k == "visuals" else val)
                       for k, val in _SAMPLE_EXTRA.get(i, {}).items()}})
    return out
