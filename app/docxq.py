"""One exam question in Word, laid out like 出題軟體: （　　）1. for choice questions, options in
4 / 2 / 1 tab-aligned columns by length, a 「右圖」 figure beside the question (a borderless
two-cell table), tables as real Word tables, SVG figures as sharp pictures, and each question
kept on one page."""
import re

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩"
LEAD_BRACKET = re.compile(r"^\s*[（(]\s*[　 ]*\s*[)）]\s*")
TEXT_W = 17.8  # cm between the 1.6 cm margins of A4


def is_choice(q: dict) -> bool:
    return bool(q.get("options")) or bool(re.search(r"選擇|是非", q.get("qtype") or ""))


def opt_len(o: str) -> int:
    o = re.sub(r"\$[^$]*\$", "xx", o or "")
    o = re.sub(r"\{([^|}]*)\|[^}]*\}", r"\1", o)
    return len(o.replace("**", ""))


def _keep(p, next_too=True):
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = next_too
    p.paragraph_format.space_after = Pt(0)


def _no_borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    borders = pr.makeelement(qn("w:tblBorders"), {})
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(borders.makeelement(qn(f"w:{side}"), {qn("w:val"): "nil"}))
    pr.append(borders)


def figure_table(doc_or_cell, t: dict, size: int, add_text, plain):
    rows = ([t.get("header")] if t.get("header") else []) + list(t.get("rows") or [])
    if not rows:
        return None
    ncol = max(len(r) for r in rows)
    tb = doc_or_cell.add_table(rows=len(rows), cols=ncol)
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tb.autofit = False
    longest = [max(len(str(r[j])) if j < len(r) else 0 for r in rows) for j in range(ncol)]
    widths = [min(5.5, max(1.4, 0.45 * L + 0.6)) for L in longest]  # as wide as the words, not the page
    for j, col in enumerate(tb.columns):
        col.width = Cm(widths[j])
    for i, r in enumerate(rows):
        for j in range(ncol):
            cell = tb.rows[i].cells[j]
            cell.width = Cm(widths[j])
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = 1
            add_text(p, plain(str(r[j]) if j < len(r) else ""), size - 1, bold=bool(t.get("header")) and i == 0)
    return tb


def pic_width(path: str, room: float) -> float:
    """How wide a figure prints (cm): wide figures (a table with four graphs) get the line,
    squarish ones stay modest — never wider than the space it is in."""
    try:
        import cv2
        h, w = cv2.imread(path).shape[:2]
        ratio = w / max(1, h)
    except Exception:
        ratio = 1.5
    want = 15.0 if ratio > 2.6 else 12.5 if ratio > 1.8 else 9.5 if ratio > 1.2 else 7.0
    return min(want, room)


def add_question(doc, q: dict, n, cfg: dict, size: int, add_text, plain, media, fig_png: str | None = None):
    choice = is_choice(q)
    num = str(q.get("number") or n)
    stem = LEAD_BRACKET.sub("", q.get("stem") or "") if choice else (q.get("stem") or "")
    show_fig = cfg.get("show_figures", True)
    pic = None
    if show_fig:
        if fig_png:
            pic = fig_png
        elif q.get("figure_path") and (media / q["figure_path"]).exists() and not q.get("figure_table"):
            pic = str(media / q["figure_path"])
    table = q.get("figure_table") if show_fig and (q.get("figure_table") or {}).get("rows") else None
    right = bool(pic) and q.get("figure_pos") == "right"
    # a whole exam printed like the paper: the figure as big as it is on the original, never bigger
    frac = max(0.18, float(q.get("fig_frac") or 0)) if q.get("fig_frac") else 1.0

    # where the question text goes: the page, or the left cell of a [text | figure] row
    if right:
        row = doc.add_table(rows=1, cols=2)
        _no_borders(row)
        left, rcell = row.rows[0].cells
        left.width, rcell.width = Cm(TEXT_W * 0.62), Cm(TEXT_W * 0.38)
        holder = left
        first = left.paragraphs[0]
    else:
        holder = doc
        first = doc.add_paragraph()
    first.paragraph_format.space_before = Pt(8)
    first.paragraph_format.left_indent = Cm(1.9 if choice else 0.9)
    first.paragraph_format.first_line_indent = Cm(-1.9 if choice else -0.9)
    _keep(first)
    add_text(first, ("（　　）" if choice else "") + f"{num}. ", size, bold=False)
    lines = re.sub(r"\\n(?![A-Za-z])", "\n", stem or "").split("\n")
    add_text(first, lines[0], size)
    for line in lines[1:]:
        p = holder.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.9 if choice else 0.9)
        _keep(p)
        add_text(p, line, size)
    if table:
        figure_table(holder, table, size, add_text, plain)
    if pic and not right:
        p = holder.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.9 if choice else 0.9)
        _keep(p)
        try:
            p.add_run().add_picture(pic, width=Cm(min(pic_width(pic, TEXT_W - 2.0), TEXT_W * frac)))
        except Exception:
            pass
    opts = q.get("options") or []
    if opts:
        longest = max(opt_len(o) for o in opts)
        room = TEXT_W * (0.55 if right else 1.0)  # options beside a 右圖 have only the left cell (less its padding)
        per = 4 if longest <= (4 if right else 7) and len(opts) >= 3 else 2 if longest <= (8 if right else 16) else 1
        indent = 1.9
        step = (room - indent) / per
        for k in range(0, len(opts), per):
            p = holder.add_paragraph()
            p.paragraph_format.left_indent = Cm(indent)
            for j in range(1, per):
                p.paragraph_format.tab_stops.add_tab_stop(Cm(indent + j * step), WD_TAB_ALIGNMENT.LEFT)
            _keep(p, k + per < len(opts))
            for j, o in enumerate(opts[k:k + per]):
                i = k + j
                if j:
                    p.add_run("\t")
                add_text(p, f"{CIRCLED[i] if i < 10 else i + 1} {o}", size)
    if right:
        rp = rcell.paragraphs[0]
        rp.alignment = 1
        try:
            rp.add_run().add_picture(pic, width=Cm(min(pic_width(pic, TEXT_W * 0.36), TEXT_W * frac)))
        except Exception:
            pass
    # the question stays together, but the next one may go to another page (a 是非題 has no option
    # line to end the keep-with-next chain: six questions once moved to the next page as one block)
    if holder.paragraphs:
        holder.paragraphs[-1].paragraph_format.keep_with_next = False
    return choice


def option_notes(doc, q: dict, size: int, add_text):
    for x in q.get("option_notes") or []:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.2)
        p.paragraph_format.space_after = Pt(0)
        add_text(p, x, size - 2, color=(60, 60, 60))
